"""The usage report as an Excel workbook (openpyxl, write-only: large reports stream).

Sheets (group 2 contract section 4): **Özet** (report details, method note, company totals, top
sites and categories, notable changes, alerts, sealing of the previous day, data quality),
**Departmanlar**, **Kişiler**, **Kişi-Site** (every person x site), **Kişi-Saat** (every person x
hour), **Engellenenler** (every blocked attempt the report lists — a list cut at the report's
limit ends with a row saying so), **Cihazlar** (every device with traffic or with blocked attempts,
also the ones without a person).

* Durations are Excel **time values** (days; ``[h]:mm:ss``, so 26 hours show as ``26:00:00``)
  and volumes are **numbers of bytes** (``#,##0``): both sort, filter and add up in Excel; no
  number is written as text.
* Times are written in site time (Europe/Istanbul) without a time zone, as Excel has none.
* Every table sheet has a styled header row that stays visible (frozen) and an auto filter.
* **Formula injection:** a text that starts with ``=``, ``+``, ``-``, ``@``, a tab or a carriage
  return gets a leading ``'`` and is always stored as text (never as a formula) — names, domains
  and categories come from the network and from people (the CSV exports use the same rule).
  Control characters an Excel file cannot hold (a DHCP host name may carry one) are written
  visibly as ``\\xNN``, as in the usage exports.
* Excel holds at most 1 048 576 rows per sheet: a longer list is cut with a last row saying so.
"""

from __future__ import annotations

import io
from collections.abc import Iterable, Mapping, Sequence
from datetime import datetime
from typing import Any, Final

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from arcilogger_core.timeutil import to_site
from arcilogger_worker.alerts.texts import SEVERITY_LABELS, STATE_LABELS
from arcilogger_worker.reports import texts as t
from arcilogger_worker.reports.html import seal_text, summary_line

DURATION: Final = "[h]:mm:ss"
BYTES: Final = "#,##0"
INTEGER: Final = "#,##0"
IDENTIFIER: Final = "0"
DATETIME: Final = "dd.mm.yyyy hh:mm:ss"
PERCENT: Final = "0%"
#: Rows of one sheet including the header (Excel's limit is 1 048 576).
MAX_ROWS: Final = 1_048_575
FORMULA_PREFIXES: Final = ("=", "+", "-", "@", "\t", "\r")
HEADER_FONT: Final = Font(bold=True, color="0F172A")
HEADER_FILL: Final = PatternFill("solid", fgColor="E2E8F0")
TITLE_FONT: Final = Font(bold=True, size=14, color="0F766E")
SECTION_FONT: Final = Font(bold=True, size=11, color="0F766E")
NOTE_FONT: Final = Font(italic=True, color="92400E")
THIN: Final = Side(style="thin", color="CBD5E1")


class Column:
    """A column of a table sheet: header, width and how a value is written."""

    def __init__(self, header: str, width: float, kind: str = "text") -> None:
        self.header = header
        self.width = width
        self.kind = kind


def safe_text(value: str) -> str:
    """A text Excel can hold and never reads as a formula (see the module text)."""
    value = ILLEGAL_CHARACTERS_RE.sub(lambda found: f"\\x{ord(found.group()):02x}", value)
    return "'" + value if value.startswith(FORMULA_PREFIXES) else value


def _local(value: Any) -> datetime | None:
    found = t.moment(value)
    return None if found is None else to_site(found).replace(tzinfo=None)


class _Sheets:
    def __init__(self) -> None:
        self.book = Workbook(write_only=True)

    def cell(self, sheet: Any, value: Any, kind: str = "text", **style: Any) -> Any:
        """A cell with the value in its Excel type (``kind``: text, int, id, bytes, duration,
        datetime, percent)."""
        if value is None or value == "":
            cell = WriteOnlyCell(sheet, value=None)
        elif kind == "duration" and isinstance(value, int | float):
            cell = WriteOnlyCell(sheet, value=value / 86400)
            cell.number_format = DURATION
        elif kind in ("bytes", "int", "id") and isinstance(value, int | float):
            cell = WriteOnlyCell(sheet, value=value)
            cell.number_format = {"bytes": BYTES, "int": INTEGER, "id": IDENTIFIER}[kind]
        elif kind == "percent" and isinstance(value, int | float):
            cell = WriteOnlyCell(sheet, value=value)
            cell.number_format = PERCENT
        elif kind == "datetime":
            moment = _local(value)
            cell = WriteOnlyCell(sheet, value=moment)
            if moment is not None:
                cell.number_format = DATETIME
        else:
            cell = WriteOnlyCell(sheet, value=safe_text(str(value)))
            cell.data_type = "s"  # never a formula, whatever the text
        for name, setting in style.items():
            setattr(cell, name, setting)
        return cell

    def table(
        self,
        title: str,
        columns: Sequence[Column],
        rows: Iterable[Sequence[Any]],
        *,
        note: str | None = None,
    ) -> int:
        """A sheet with a frozen, filtered header row (``note``: a last row under the table);
        returns the number of data rows."""
        sheet = self.book.create_sheet(title)
        sheet.freeze_panes = "A2"
        for index, column in enumerate(columns, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = column.width
        sheet.append(
            [
                self.cell(
                    sheet,
                    column.header,
                    font=HEADER_FONT,
                    fill=HEADER_FILL,
                    border=Border(bottom=THIN),
                    alignment=Alignment(wrap_text=True, vertical="top"),
                )
                for column in columns
            ]
        )
        count = 0
        for row in rows:
            if count + 1 >= MAX_ROWS:
                sheet.append(
                    [self.cell(sheet, "Excel satır sınırı: liste burada kesildi.", font=NOTE_FONT)]
                )
                break
            sheet.append(
                [
                    self.cell(sheet, value, column.kind)
                    for value, column in zip(row, columns, strict=True)
                ]
            )
            count += 1
        if note is not None:
            sheet.append([self.cell(sheet, note, font=NOTE_FONT)])
        last = get_column_letter(len(columns))
        sheet.auto_filter.ref = f"A1:{last}{count + 1}"
        return count


def _summary(sheets: _Sheets, document: Mapping[str, Any]) -> None:
    report, data = document["report"], document["data"]
    company = data.get("company") or {}
    sheet = sheets.book.create_sheet("Özet")
    sheet.column_dimensions["A"].width = 44
    sheet.column_dimensions["B"].width = 26
    sheet.column_dimensions["C"].width = 16
    sheet.column_dimensions["D"].width = 16
    sheet.column_dimensions["E"].width = 16
    sheet.freeze_panes = "A2"
    cell = sheets.cell

    def line(*values: Any) -> None:
        sheet.append(list(values))

    def section(text: str) -> None:
        line()
        line(cell(sheet, text, font=SECTION_FONT))

    def header(*names: str) -> None:
        line(*(cell(sheet, name, font=HEADER_FONT, fill=HEADER_FILL) for name in names))

    line(cell(sheet, f"{t.PRODUCT} kullanım raporu — {report['name']}", font=TITLE_FONT))
    line(cell(sheet, summary_line(document, 0)))
    line(cell(sheet, summary_line(document, 1)))
    line(cell(sheet, t.METHOD_NOTE_TR, font=NOTE_FONT))
    line(cell(sheet, t.METHOD_NOTE_EN, font=NOTE_FONT))
    section("Rapor / Report")
    site = t.site_name(report)
    for label, value, kind in (
        ("Rapor", report["name"], "text"),
        ("Tür", t.kind_name(report["kind"]), "text"),
        ("Dönem başlangıcı", data["period"]["from"], "datetime"),
        ("Dönem sonu (hariç)", data["period"]["to"], "datetime"),
        ("Oluşturma zamanı", report.get("generated_at"), "datetime"),
        ("Hesaplanan son an", data.get("computed_until"), "datetime"),
        ("Kişi düzeyinde", "evet" if data.get("person_level") else "hayır", "text"),
        ("Takma adlarla", "evet" if data.get("pseudonymized") else "hayır", "text"),
        # The organisation; no row at all when no name is set.
        *((("Site", site, "text"),) if site else ()),
        ("Saat dilimi", "Europe/Istanbul (UTC+3)", "text"),
    ):
        line(cell(sheet, label), cell(sheet, value, kind))
    section("Şirket geneli / Company totals (süreler tahmini)")
    for label, key, kind in (
        ("Aktif kullanıcı", "active_users", "int"),
        ("Çevrimiçi süre (toplam)", "online_seconds", "duration"),
        ("Aktif süre (toplam)", "active_seconds", "duration"),
        ("Mesai dışı aktif süre", "off_hours_active_seconds", "duration"),
        ("İndirilen (bayt)", "download_bytes", "bytes"),
        ("Yüklenen (bayt)", "upload_bytes", "bytes"),
        ("Oturum", "sessions", "int"),
        ("Engellenen deneme", "blocked", "int"),
    ):
        line(cell(sheet, label), cell(sheet, company.get(key), kind))
    section("En çok kullanılan siteler / Top sites")
    header("Site", "Kategori", "Süre", "Pay", "İndirilen (bayt)")
    for site in company.get("top_sites") or []:
        line(
            cell(sheet, site.get("domain")),
            cell(sheet, site.get("category") or t.UNKNOWN_CATEGORY),
            cell(sheet, site.get("seconds"), "duration"),
            cell(sheet, site.get("share"), "percent"),
            cell(sheet, site.get("download_bytes"), "bytes"),
        )
    section("En çok kullanılan kategoriler / Top categories")
    header("Kategori", "Süre", "Pay")
    for item in company.get("top_categories") or []:
        line(
            cell(sheet, item.get("category") or t.UNKNOWN_CATEGORY),
            cell(sheet, item.get("seconds"), "duration"),
            cell(sheet, item.get("share"), "percent"),
        )
    section("Dikkat çeken değişiklikler / Notable changes (7 günlük ortalamaya göre)")
    header("Değişiklik", "Kapsam", "Değişim")
    for item in data.get("changes") or []:
        line(
            cell(sheet, item.get("text_tr")),
            cell(sheet, item.get("scope")),
            cell(sheet, (item.get("change_percent") or 0) / 100, "percent"),
        )
    alerts = document.get("alerts") or {}
    section(f"Güvenlik uyarıları / Security alerts ({int(alerts.get('total') or 0)})")
    header("Zaman", "Önem", "Kural", "Uyarı", "Durum")
    for item in alerts.get("listed") or []:
        severity, state = str(item.get("severity")), str(item.get("state"))
        line(
            cell(sheet, item.get("time"), "datetime"),
            cell(sheet, SEVERITY_LABELS.get(severity, (severity,))[0]),
            cell(sheet, item.get("rule")),
            cell(sheet, f"{item.get('name_tr')}: {item.get('text_tr')}"),
            cell(sheet, STATE_LABELS.get(state, (state,))[0]),
        )
    seal = document.get("seal") or {}
    if seal:
        section("Önceki günün mühürü / Sealing of the previous day")
        line(cell(sheet, seal_text(seal, 0)))
        line(cell(sheet, seal_text(seal, 1)))
    quality = data.get("data_quality") or {}
    section("Veri kalitesi / Data quality")
    line(
        cell(sheet, "Kişiye bağlanamayan süre"),
        cell(sheet, quality.get("unattributed_seconds"), "duration"),
    )
    for note in [*(quality.get("notes_tr") or []), *(quality.get("notes_en") or [])]:
        line(cell(sheet, note))


def _blocked_note(data: Mapping[str, Any]) -> str | None:
    """The last row of a cut list of blocked attempts (a report lists at most 50 000)."""
    listed = len(data.get("blocked") or [])
    total = data.get("blocked_total")
    if not isinstance(total, int) or total <= listed:
        return None
    return (
        f"Liste kesildi: {listed} / {total} engellenen deneme listelendi; kişi sayıları "
        "listelenmeyenleri de sayar. / The list was cut; the counts include the ones not listed."
    )


def render_xlsx(document: Mapping[str, Any]) -> bytes:
    """The workbook of ``document``."""
    data = document["data"]
    sheets = _Sheets()
    sheets.book.properties.creator = t.PRODUCT
    sheets.book.properties.title = f"{t.PRODUCT} kullanım raporu — {document['report']['name']}"
    generated = _local(document["report"].get("generated_at"))
    if generated is not None:
        sheets.book.properties.created = generated
        sheets.book.properties.modified = generated
    _summary(sheets, document)
    sheets.table(
        "Departmanlar",
        [
            Column("Departman", 28),
            Column("Aktif kullanıcı", 14, "int"),
            Column("Aktif süre", 14, "duration"),
            Column("İndirilen (bayt)", 18, "bytes"),
            Column("Yüklenen (bayt)", 18, "bytes"),
            Column("Engellenen", 12, "int"),
        ],
        (
            [
                t.department(row.get("name")),
                row.get("active_users"),
                row.get("active_seconds"),
                row.get("download_bytes"),
                row.get("upload_bytes"),
                row.get("blocked"),
            ]
            for row in data.get("departments") or []
        ),
    )
    sheets.table(
        "Kişiler",
        [
            Column("Kişi", 26),
            Column("Departman", 22),
            Column("İlk etkinlik", 19, "datetime"),
            Column("Son etkinlik", 19, "datetime"),
            Column("Çevrimiçi süre", 14, "duration"),
            Column("Aktif süre", 14, "duration"),
            Column("Mesai dışı aktif süre", 14, "duration"),
            Column("İndirilen (bayt)", 18, "bytes"),
            Column("Yüklenen (bayt)", 18, "bytes"),
            Column("Oturum", 10, "int"),
            Column("Engellenen", 11, "int"),
            Column("En çok kullanılan siteler", 60),
            Column("En çok kategoriler", 50),
        ],
        (
            [
                person.get("label"),
                t.department(person.get("department")),
                person.get("first_activity"),
                person.get("last_activity"),
                person.get("online_seconds"),
                person.get("active_seconds"),
                person.get("off_hours_active_seconds"),
                person.get("download_bytes"),
                person.get("upload_bytes"),
                person.get("sessions"),
                person.get("blocked"),
                ", ".join(
                    f"{site.get('domain')} ({t.duration(site.get('seconds'))})"
                    for site in person.get("top_sites") or []
                ),
                ", ".join(
                    f"{item.get('category') or t.UNKNOWN_CATEGORY} "
                    f"({t.duration(item.get('seconds'))})"
                    for item in person.get("top_categories") or []
                ),
            ]
            for person in data.get("persons") or []
        ),
    )
    sheets.table(
        "Kişi-Site",
        [
            Column("Kişi", 26),
            Column("Departman", 22),
            Column("Site", 34),
            Column("Kategori", 30),
            Column("Süre", 12, "duration"),
            Column("İndirilen (bayt)", 18, "bytes"),
            Column("Yüklenen (bayt)", 18, "bytes"),
            Column("Oturum", 10, "int"),
        ],
        (
            [
                row.get("label"),
                t.department(row.get("department")),
                row.get("domain"),
                row.get("category") or t.UNKNOWN_CATEGORY,
                row.get("seconds"),
                row.get("download_bytes"),
                row.get("upload_bytes"),
                row.get("sessions"),
            ]
            for row in data.get("person_sites") or []
        ),
    )
    sheets.table(
        "Kişi-Saat",
        [
            Column("Kişi", 26),
            Column("Saat", 19, "datetime"),
            Column("Aktif süre", 12, "duration"),
            Column("İndirilen (bayt)", 18, "bytes"),
            Column("Yüklenen (bayt)", 18, "bytes"),
        ],
        (
            [
                row.get("label"),
                row.get("hour"),
                row.get("active_seconds"),
                row.get("download_bytes"),
                row.get("upload_bytes"),
            ]
            for row in data.get("person_hours") or []
        ),
    )
    sheets.table(
        "Engellenenler",
        [
            Column("Zaman", 19, "datetime"),
            Column("Kişi", 26),
            Column("Site", 34),
            Column("Kategori", 30),
            Column("İşlem", 12),
            Column("Kaynak", 20),
        ],
        (
            [
                row.get("time"),
                row.get("label"),
                row.get("domain"),
                row.get("category") or t.UNKNOWN_CATEGORY,
                row.get("action"),
                t.SOURCE_LABELS.get(str(row.get("source")), (row.get("source"),))[0],
            ]
            for row in data.get("blocked") or []
        ),
        note=_blocked_note(data),
    )
    sheets.table(
        "Cihazlar",
        [
            Column("Cihaz", 24),
            Column("MAC", 19),
            Column("IP", 16),
            Column("VLAN", 8, "id"),
            Column("Kişi", 26),
            Column("Aktif süre", 12, "duration"),
            Column("İndirilen (bayt)", 18, "bytes"),
            Column("Yüklenen (bayt)", 18, "bytes"),
        ],
        (
            [
                row.get("label"),
                row.get("mac"),
                row.get("ip"),
                row.get("vlan"),
                row.get("person") or "—",
                row.get("active_seconds"),
                row.get("download_bytes"),
                row.get("upload_bytes"),
            ]
            for row in data.get("devices") or []
        ),
    )
    buffer = io.BytesIO()
    sheets.book.save(buffer)
    return buffer.getvalue()
