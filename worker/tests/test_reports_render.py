"""The three renderings of a usage report — HTML e-mail, PDF, Excel — against golden files made
from the fixed sample (:mod:`arcilogger_worker.reports.testing`), plus the rules each one must
keep: e-mail-safe HTML with no external resources, Turkish text in the PDF, typed and
formula-safe cells in the workbook, person-level detail with names by default.

Golden files: ``worker/tests/snapshots/usage-report.*``. After an intended change, write them
again with ``ARCILOGGER_UPDATE_SNAPSHOTS=1`` and review the diff.
"""

from __future__ import annotations

import io
import json
import os
import re
from collections.abc import Callable
from datetime import timedelta
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from arcilogger_worker.pdf.text import normalized, pdf_text
from arcilogger_worker.reports.content import ReportMeta, assemble
from arcilogger_worker.reports.html import render_html
from arcilogger_worker.reports.mail import delivery_lines, render_mail, subject
from arcilogger_worker.reports.pdf import render_pdf
from arcilogger_worker.reports.testing import GENERATED_AT, SITE_NAME, report_sample
from arcilogger_worker.reports.texts import METHOD_NOTE_EN, METHOD_NOTE_TR, NO_DEPARTMENT
from arcilogger_worker.reports.xlsx import render_xlsx, safe_text
from arcilogger_worker.usage.query import UNKNOWN_DEPARTMENT
from arcilogger_worker.usage.report_data import sample_report_data

SNAPSHOTS = Path(__file__).resolve().parent / "snapshots"
UPDATE = os.environ.get("ARCILOGGER_UPDATE_SNAPSHOTS") == "1"
WEB = "https://172.20.20.15"
ALERTS: dict[str, Any] = {
    "total": 2,
    "listed": [
        {
            "time": "2026-10-19T11:58:01+03:00",
            "rule": "R15",
            "severity": "high",
            "state": "resolved",
            "name_tr": "Log kaynağı sessiz",
            "name_en": "Log source silent",
            "text_tr": "fortigate 10 dakikadan uzun süredir log göndermiyor.",
            "text_en": "fortigate sent no log for more than 10 minutes.",
        },
        {
            "time": "2026-10-19T03:05:00+03:00",
            "rule": "R20-disk",
            "severity": "medium",
            "state": "open",
            "name_tr": "Disk dolmak üzere",
            "name_en": "Disk almost full",
            "text_tr": "D: %83 dolu (eşik %80).",
            "text_en": "D: is 83 % full (threshold 80 %).",
        },
    ],
    "by_severity": {"high": 1, "medium": 1},
}
SEAL: dict[str, Any] = {
    "day": "2026-10-19",
    "status": "sealed",
    "sealed_at": "2026-10-20T00:31:12+03:00",
    "serial": "4F2A19",
    "provider": "zamane",
    "test_stamp": False,
}


def make_document(data: dict[str, Any] | None = None, **meta: Any) -> dict[str, Any]:
    document = assemble(
        report_sample() if data is None else data,
        meta=ReportMeta(
            name=meta.pop("name", "Tam gün raporu"),
            kind=meta.pop("kind", "full_day"),
            site_name=meta.pop("site_name", SITE_NAME),
            web_url=WEB,
            generated_at=GENERATED_AT,
            run_id=meta.pop("run_id", 17),
            schedule_id=2,
            **meta,
        ),
        alerts=ALERTS,
        seal=SEAL,
    )
    document["report"]["version"] = "0.0.0-test"  # the golden files outlive releases
    return document


def golden(name: str, actual: str) -> None:
    path = SNAPSHOTS / name
    if UPDATE:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(actual, encoding="utf-8", newline="\n")
    assert path.is_file(), f"missing golden file {name}: run with ARCILOGGER_UPDATE_SNAPSHOTS=1"
    expected = path.read_text(encoding="utf-8")
    assert actual == expected, (
        f"{name} changed; review and update with ARCILOGGER_UPDATE_SNAPSHOTS=1"
    )


def persons(data: dict[str, Any]) -> list[str]:
    return [str(person["label"]) for person in data["persons"]]


# -- HTML e-mail ----------------------------------------------------------------------------------


def email_html() -> str:
    document = make_document()
    html, _ = render_mail(
        document,
        email_html=True,
        attachments=[
            {"name": "ARCILOGGER-kullanim-tam-gun-20261019-r17.pdf", "size": 95_000},
            {"name": "ARCILOGGER-kullanim-tam-gun-20261019-r17.xlsx", "size": 41_000},
        ],
        omitted=False,
        limit_mb=15,
        run_id=17,
    )
    return html


def test_email_html_matches_the_golden_file() -> None:
    golden("usage-report.email.html", email_html())


def test_email_html_is_safe_for_mail_clients_and_loads_nothing() -> None:
    html = email_html()
    lowered = html.lower()
    for forbidden in ("<script", "<style", "<img", "<link", "<iframe", "<form", "@import"):
        assert forbidden not in lowered, forbidden
    assert "url(" not in lowered
    assert "javascript:" not in lowered
    assert " class=" not in lowered  # inline styles only (Gmail drops <style>)
    styles = re.findall(r'style="([^"]*)"', html)
    assert styles
    for style in styles:
        assert "flex" not in style, style
        assert "grid" not in style, style
        assert "position:" not in style, style
    # The only addresses: the ARCILOGGER web interface.
    links = re.findall(r"""(?:href|src)\s*=\s*["']([^"']*)["']""", html, flags=re.IGNORECASE)
    assert links == [f"{WEB}/reports"]
    for url in re.findall(r"(?:https?|ftp)://[^\s\"'<>)]+", html):
        assert url.startswith(WEB + "/"), url
    assert html.count('role="presentation"') > 5  # table layout, widths as attributes
    assert 'bgcolor="' in html
    assert 'width="680"' in html
    # Turkish first, an English summary line, the method note in both languages.
    assert '<html lang="tr">' in html
    assert "Full-day report 19.10.2026 (Örnek A.Ş.): 18 active users" in html
    assert METHOD_NOTE_TR in html
    assert METHOD_NOTE_EN in html


def test_email_html_lists_every_person_and_every_section() -> None:
    html = email_html()
    data = report_sample()
    for name in persons(data):
        assert f"<b>{name}</b>" in html, name
    for heading in (
        "Şirket geneli",
        "Departmanlar",
        "Kişiler (18)",
        "Dikkat çeken değişiklikler",
        "Güvenlik uyarıları (2)",
        "Önceki günün mühürü",
        "Veri kalitesi",
        "Ekler",
    ):
        assert heading in html, heading
    assert "19.10.2026: Mühürlendi (20.10.2026 00:31, seri 4F2A19)" in html
    assert "Video ve müzik akışı 7 günlük ortalamanın %180 üzerinde." in html
    assert "Log kaynağı sessiz" in html
    assert "Disk dolmak üzere" in html
    assert "ARCILOGGER-kullanim-tam-gun-20261019-r17.pdf (92,8 KB)" in html


def test_values_from_the_network_are_escaped() -> None:
    data = report_sample()
    data["persons"][0]["label"] = '<script>alert("x")</script>'
    data["persons"][0]["top_sites"][0]["domain"] = 'evil.example"><img src=x>'
    html = render_html(make_document(data))
    assert "<script>" not in html
    assert "<img" not in html
    assert "&lt;script&gt;alert(&quot;x&quot;)&lt;/script&gt;" in html


def test_an_oversized_report_links_to_the_download_instead() -> None:
    lines = delivery_lines(
        [{"name": "a.pdf", "size": 20 * 1024 * 1024}],
        omitted=True,
        limit_mb=15,
        link=f"{WEB}/reports?run=17",
    )
    assert "Ekler e-postaya eklenmedi: toplam 20 MB, sınır 15 MB" in lines[0]
    assert lines[0].endswith(f"{WEB}/reports?run=17")
    assert lines[1].startswith("The files were not attached")


def test_short_mail_without_the_html_report() -> None:
    document = make_document()
    html, text = render_mail(
        document, email_html=False, attachments=[], omitted=False, limit_mb=15, run_id=17
    )
    assert "Kişiler" not in html  # not the report itself
    assert METHOD_NOTE_TR in html
    assert METHOD_NOTE_TR in text
    assert f"{WEB}/reports?run=17" in text
    assert text.startswith("ARCILOGGER kullanım raporu — Tam gün raporu\n")
    assert subject(document) == "[ARCILOGGER] Tam gün raporu — 19.10.2026 (Örnek A.Ş.)"


# -- PDF ------------------------------------------------------------------------------------------


def test_pdf_text_matches_the_golden_file() -> None:
    rendered = render_pdf(make_document())
    golden("usage-report.pdf.txt", pdf_text(rendered.pdf))


def test_pdf_has_the_sections_page_numbers_and_turkish_text() -> None:
    data = report_sample()
    rendered = render_pdf(make_document())
    text = normalized(pdf_text(rendered.pdf))
    assert rendered.pages >= 1 + len(data["departments"])  # cover + a page per department
    assert f"Sayfa 1 / {rendered.pages}" in text
    assert f"Sayfa {rendered.pages} / {rendered.pages}" in text
    assert "Oluşturma 20.10.2026 07:00" in text
    assert METHOD_NOTE_TR in text
    assert METHOD_NOTE_EN in text
    for department in data["departments"]:
        assert f"Departman: {department['name']}" in text, department["name"]
    for name in persons(data):
        assert name in text, name
    for letter in "çğıöşüÇİÖŞÜ":  # Turkish letters survive (ToUnicode maps of the subsets)
        assert letter in text, letter
    assert b"Helvetica" not in rendered.pdf  # only the embedded DejaVu subsets
    # Per person: top sites with time, categories, blocked attempts with their source.
    assert "Engellenen denemeler (5) / Blocked attempts" in text
    assert "www.bet365.com Gambling Web filtresi" in text


def _with_subjects(data: dict[str, Any]) -> dict[str, Any]:
    """Subject keys on the persons and their detail rows, as the engine sends them."""
    keys = {person["label"]: f"person:{index}" for index, person in enumerate(data["persons"], 1)}
    for name in ("persons", "person_sites", "blocked"):
        for row in data[name]:
            row["subject"] = keys[row["label"]]
    return data


def test_persons_who_share_a_name_keep_their_own_blocked_attempts() -> None:
    """``persons.name`` is not unique and a device may carry a person's name: the PDF lists
    each person's attempts by the subject, never by the name shown."""
    data = _with_subjects(report_sample())
    for name in ("persons", "person_sites", "blocked"):
        for row in data[name]:
            if row["label"] == "Deniz Koç":
                row["label"] = "Ece Doğan"  # two employees with one name
    data["blocked"].append(
        {
            "time": "2026-10-19T16:00:00+03:00",
            "subject": "device:00:0c:29:00:17:01",
            "label": "Ece Doğan",  # a device named like a person
            "domain": "yalniz-cihaz.example",
            "category": "Gambling",
            "action": "blocked",
            "source": "web",
        }
    )
    text = normalized(pdf_text(render_pdf(make_document(data)).pdf))
    assert "Engellenen denemeler (5) / Blocked attempts" in text  # Ece Doğan
    assert "Engellenen denemeler (3) / Blocked attempts" in text  # the other one
    assert "Engellenen denemeler (8)" not in text
    assert "Engellenen denemeler (9)" not in text
    for domain in ("free-proxy.example.org", "torrent.example.net"):
        listed = sum(1 for row in data["blocked"] if row["domain"] == domain)
        assert text.count(domain) == listed, domain  # each attempt under one person only
    assert "yalniz-cihaz.example" not in text  # the device's attempt: the workbook only


def test_a_weekly_report_dates_the_first_and_last_activity() -> None:
    """Over several days the first and the last activity may fall on different days."""
    data = report_sample()
    data["kind"] = "weekly"
    data["period"] = {"from": "2026-10-12T00:00:00+03:00", "to": "2026-10-19T00:00:00+03:00"}
    person = data["persons"][0]
    person["first_activity"] = "2026-10-12T08:05:00+03:00"
    person["last_activity"] = "2026-10-16T17:58:00+03:00"
    document = make_document(data, kind="weekly", name="Haftalık rapor")
    text = normalized(pdf_text(render_pdf(document).pdf))
    assert "12.10.2026 08:05 - 16.10.2026 17:58" in text
    html = render_html(document)
    assert "12.10.2026 08:05<br>16.10.2026 17:58" in html
    # One day: the clock times only (the date is the report's).
    one_day = render_html(make_document())
    first = report_sample()["persons"][0]
    assert f"{first['first_activity'][11:16]}-{first['last_activity'][11:16]}" in one_day


def test_a_cut_list_of_blocked_attempts_says_so() -> None:
    """A report lists at most 50 000 attempts (and the engine keeps at most
    ``usage.blocked_per_day`` per subject and day): the count can exceed the list."""
    data = report_sample()
    ece = next(person for person in data["persons"] if person["label"] == "Ece Doğan")
    ece["blocked"] = 7  # 5 listed
    data["blocked_total"] = len(data["blocked"]) + 1000
    document = make_document(data)
    text = normalized(pdf_text(render_pdf(document).pdf))
    assert "Engellenen denemeler (7) / Blocked attempts" in text
    assert "… 2 deneme daha sayıldı ama listelenmedi" in text
    sheet = load_workbook(io.BytesIO(render_xlsx(document)))["Engellenenler"]
    assert sheet.max_row == len(data["blocked"]) + 2  # header, the rows, the note
    note = str(sheet.cell(sheet.max_row, 1).value)
    assert note.startswith(f"Liste kesildi: {len(data['blocked'])} / {len(data['blocked']) + 1000}")
    assert sheet.auto_filter.ref == f"A1:F{len(data['blocked']) + 1}"


def test_the_group_without_a_department_has_one_name() -> None:
    """The engine names it "Bilinmiyor" (its department rows include the devices and addresses
    without one); the persons of that group carry no department."""
    assert NO_DEPARTMENT[0] == UNKNOWN_DEPARTMENT
    data = report_sample()
    assert UNKNOWN_DEPARTMENT in [row["name"] for row in data["departments"]]
    document = make_document(data)
    text = normalized(pdf_text(render_pdf(document).pdf))
    assert "Departmanı yok" not in text
    page = text[text.index(f"Departman: {UNKNOWN_DEPARTMENT}") :]
    assert page.index("Aktif kullanıcı / Active users") < page.index("Kişiler (1)")
    html = render_html(document)
    assert "Departmanı yok" not in html
    assert "<b>Kerem Tunç</b><br>" in html
    book = load_workbook(io.BytesIO(render_xlsx(document)))
    rows = book["Kişiler"].iter_rows(values_only=True)
    assert next(row for row in rows if row[0] == "Kerem Tunç")[1] == UNKNOWN_DEPARTMENT


def test_without_an_organisation_name_no_rendering_leaves_a_gap() -> None:
    """``[site] site_name`` is empty until the installer sets it: the subject, the summary
    sentences, the HTML header, the PDF cover and footer leave the name out, the workbook has
    no "Site" row — no empty brackets, no dangling separator."""
    for empty in ("", "   "):
        document = make_document(site_name=empty)
        assert subject(document) == "[ARCILOGGER] Tam gün raporu — 19.10.2026"
        html, text = render_mail(
            document, email_html=True, attachments=[], omitted=False, limit_mb=15, run_id=17
        )
        pdf = normalized(pdf_text(render_pdf(document).pdf))
        for rendered in (html, text, pdf):
            assert "Tam gün raporu 19.10.2026: 18 aktif kullanıcı" in rendered
            assert "Full-day report 19.10.2026: 18 active users" in rendered
            assert "()" not in rendered
            assert "( )" not in rendered
            assert SITE_NAME not in rendered
        assert 'letter-spacing:1px;">ARCILOGGER</div>' in html  # the header: the product only
        assert "· </div>" not in html
        assert "Tam gün raporu · 19.10.2026 / Full-day report" in pdf
        assert "ARCILOGGER 0.0.0-test · Oluşturma 20.10.2026 07:00" in pdf
        assert "· ·" not in pdf
        book = load_workbook(io.BytesIO(render_xlsx(document)))
        labels = [row[0] for row in book["Özet"].iter_rows(values_only=True)]
        assert "Site" not in labels[: labels.index("Saat dilimi")]
        assert labels[labels.index("Saat dilimi") - 1] == "Takma adlarla"
    # With a name, it is in each of them.
    named = make_document()
    html = render_html(named)
    pdf = normalized(pdf_text(render_pdf(named).pdf))
    assert subject(named) == f"[ARCILOGGER] Tam gün raporu — 19.10.2026 ({SITE_NAME})"
    assert f'letter-spacing:1px;">ARCILOGGER · {SITE_NAME}</div>' in html
    assert f"Tam gün raporu 19.10.2026 ({SITE_NAME}): 18 aktif kullanıcı" in html
    assert f"Tam gün raporu · 19.10.2026 · {SITE_NAME} / Full-day report" in pdf
    assert f"ARCILOGGER 0.0.0-test · {SITE_NAME} · Oluşturma" in pdf
    book = load_workbook(io.BytesIO(render_xlsx(named)))
    assert ("Site", SITE_NAME) in [row[:2] for row in book["Özet"].iter_rows(values_only=True)]


def test_pdf_is_deterministic() -> None:
    first = render_pdf(make_document())
    assert render_pdf(make_document()).pdf == first.pdf
    other = make_document()
    other["data"]["persons"][0]["label"] = "Başka Kişi"
    assert render_pdf(other).pdf != first.pdf


# -- Excel ----------------------------------------------------------------------------------------


def _cell(cell: Any) -> Any:
    value = cell.value
    if isinstance(value, timedelta):
        value = f"timedelta:{int(value.total_seconds())}"
    elif value is not None and not isinstance(value, str | int | float | bool):
        value = f"{type(value).__name__}:{value.isoformat()}"
    return [value, cell.number_format] if cell.number_format != "General" else value


def workbook_digest(data: bytes) -> dict[str, Any]:
    """Every sheet: frozen pane, filter, and its rows (tables: header, first rows, last row)."""
    book = load_workbook(io.BytesIO(data))
    digest: dict[str, Any] = {}
    for sheet in book.worksheets:
        rows = [[_cell(cell) for cell in row] for row in sheet.iter_rows()]
        if sheet.title != "Özet" and len(rows) > 8:
            rows = [*rows[:6], ["…"], rows[-1]]
        digest[sheet.title] = {
            "freeze": sheet.freeze_panes,
            "filter": sheet.auto_filter.ref,
            "rows": sheet.max_row,
            "cells": rows,
        }
    return digest


def test_xlsx_matches_the_golden_file() -> None:
    digest = workbook_digest(render_xlsx(make_document()))
    golden("usage-report.xlsx.json", json.dumps(digest, ensure_ascii=False, indent=1) + "\n")


def test_xlsx_sheets_are_typed_frozen_and_filtered() -> None:
    data = report_sample()
    book = load_workbook(io.BytesIO(render_xlsx(make_document())))
    assert book.sheetnames == [
        "Özet",
        "Departmanlar",
        "Kişiler",
        "Kişi-Site",
        "Kişi-Saat",
        "Engellenenler",
        "Cihazlar",
    ]
    expected_rows = {
        "Departmanlar": len(data["departments"]),
        "Kişiler": len(data["persons"]),
        "Kişi-Site": len(data["person_sites"]),
        "Kişi-Saat": len(data["person_hours"]),
        "Engellenenler": len(data["blocked"]),
        "Cihazlar": len(data["devices"]),
    }
    for title, count in expected_rows.items():
        sheet = book[title]
        assert sheet.max_row == count + 1, title
        assert sheet.freeze_panes == "A2", title
        assert (
            sheet.auto_filter.ref
            == f"A1:{sheet.cell(1, sheet.max_column).column_letter}{count + 1}"
        )
        assert sheet.cell(1, 1).font.bold, title
    people = book["Kişiler"]
    headers = [cell.value for cell in people[1]]
    first = dict(zip(headers, people[2], strict=True))
    person = data["persons"][0]
    # Durations are time values (days) with an [h]:mm:ss format, never text.
    active = first["Aktif süre"]
    assert active.number_format == "[h]:mm:ss"
    assert active.value == timedelta(seconds=person["active_seconds"])  # read back as time
    assert active.data_type != "s"
    # Bytes are numbers with a readable format.
    down = first["İndirilen (bayt)"]
    assert (down.value, down.number_format, down.data_type) == (
        person["download_bytes"],
        "#,##0",
        "n",
    )
    assert first["İlk etkinlik"].number_format == "dd.mm.yyyy hh:mm:ss"
    assert first["Kişi"].value == person["label"]
    # Every blocked attempt and every device (also the ones without a person).
    devices = [row[4].value for row in book["Cihazlar"].iter_rows(min_row=2)]
    assert devices.count("—") == 2
    assert book["Cihazlar"].cell(2, 4).number_format == "0"  # VLAN ids without separators


@pytest.mark.parametrize(
    "hostile",
    ['=HYPERLINK("http://evil.example","x")', "+cmd|' /C calc'!A0", "-2+3", "@SUM(1+1)", "\t=1"],
)
def test_xlsx_never_writes_a_formula(hostile: str) -> None:
    data = report_sample()
    data["persons"][0]["label"] = hostile
    data["person_sites"][0]["domain"] = hostile
    data["blocked"][0]["category"] = hostile
    book = load_workbook(io.BytesIO(render_xlsx(make_document(data))))
    found = 0
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                assert cell.data_type != "f", (sheet.title, cell.coordinate, cell.value)
                if cell.value == "'" + hostile:
                    found += 1
                    assert cell.data_type == "s"
    assert found >= 3
    assert safe_text("Elif Kaya") == "Elif Kaya"


def test_xlsx_writes_control_characters_visibly() -> None:
    """An Excel file cannot hold most control characters, and host names (DHCP, DNS) come from
    the network: one such name must not cost the report its workbook (nor its PDF and e-mail,
    rendered with it)."""
    data = report_sample()
    data["devices"][0]["label"] = "phone\x1bx"
    data["person_sites"][0]["domain"] = "evil\x01.example.com"
    data["blocked"][0]["domain"] = "=\x02cmd"
    book = load_workbook(io.BytesIO(render_xlsx(make_document(data))))
    assert book["Cihazlar"].cell(2, 1).value == "phone\\x1bx"
    assert book["Kişi-Site"].cell(2, 3).value == "evil\\x01.example.com"
    assert book["Engellenenler"].cell(2, 3).value == "'=\\x02cmd"  # and never a formula
    assert safe_text("bir\tiki\nüç") == "bir\tiki\nüç"  # tabs and line breaks are kept


# -- options --------------------------------------------------------------------------------------


Renderer = Callable[[dict[str, Any]], Any]


def test_without_person_level_only_totals_are_shown() -> None:
    data = report_sample()
    data["person_level"] = False
    for key in ("persons", "person_sites", "person_hours", "blocked"):
        data[key] = []
    document = make_document(data)
    html = render_html(document)
    assert "Bu zamanlama kişi düzeyinde değil" in html
    assert "Elif Kaya" not in html
    pdf = render_pdf(document)
    text = normalized(pdf_text(pdf.pdf))
    assert "Departman: " not in text
    assert "Elif Kaya" not in text
    book = load_workbook(io.BytesIO(render_xlsx(document)))
    assert book["Kişiler"].max_row == 1


def test_pseudonyms_and_the_preview_are_marked() -> None:
    data = report_sample()
    data["pseudonymized"] = True
    for person in data["persons"]:
        person["label"] = "Kullanıcı-7F3A"
    document = make_document(data, preview=True)
    html = render_html(document)
    assert "takma adlarla / pseudonymised" in html
    assert "Önizleme: kaydedilmedi, gönderilmedi." in html
    text = normalized(pdf_text(render_pdf(document).pdf))
    assert "Takma adlarla / Pseudonymised evet" in text
    assert "Önizleme / Preview kaydedilmedi, gönderilmedi" in text


def test_times_given_as_datetime_objects_still_render() -> None:
    """The engine's contract asks for ISO-8601 text; a ``datetime`` must not fail a report."""
    data = report_sample()
    data["persons"][0]["first_activity"] = GENERATED_AT - timedelta(hours=23)
    document = make_document(data)
    assert document["data"]["persons"][0]["first_activity"] == "2026-10-19T08:00:04+03:00"
    with pytest.raises(TypeError, match="not JSON: set"):
        make_document({**report_sample(), "extra": {1, 2}})


def test_the_engines_own_sample_renders_in_every_format() -> None:
    """``sample_report_data()`` of the usage engine (area U's at merge) satisfies the contract
    the renderers rely on: every person appears in each rendering."""
    data = sample_report_data()
    assert data["schema"] == 1
    document = make_document(data, kind=data["kind"])
    html = render_html(document)
    text = normalized(pdf_text(render_pdf(document).pdf))
    book = load_workbook(io.BytesIO(render_xlsx(document)))
    people = {cell.value for cell in book["Kişiler"]["A"][1:]}
    for name in persons(data):
        assert name in html, name
        assert name in text, name
        assert name in people, name
