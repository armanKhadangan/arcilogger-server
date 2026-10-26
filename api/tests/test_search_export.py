"""Export of a search (contract 2.12, R17, R18, 3.4): CSV and Excel, formula neutralisation, text
cells, the "Bilgi" sheet, truncation by rows and by time (name, header, sheet, audit), the
``Repr-Digest`` and audited SHA-256, re-authentication, one export at a time and the hourly
limit counted from the audit log (it survives a new application)."""

from __future__ import annotations

import base64
import csv
import hashlib
import io
from datetime import UTC, datetime, timedelta
from typing import Any

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import Engine, select
from sqlalchemy.orm import Session, sessionmaker

from arcilogger_api.app import create_app
from arcilogger_api.routers.search import state_of
from arcilogger_api.search.export import NOT_SEALED
from arcilogger_api.search.testing import Lines, sign_in
from arcilogger_api.testing import LoggedIn
from arcilogger_core.config import Config
from arcilogger_db.models import AuditLog
from arcilogger_db.settings import set_setting

API = "/api/v1/search"
NOW = datetime.now(UTC).replace(microsecond=0)
START, END = NOW - timedelta(hours=6), NOW


def setting(factory: sessionmaker[Session], key: str, value: Any) -> None:
    with factory() as session:
        set_setting(session, key, value, updated_by="test")
        session.commit()


@pytest.fixture
def exporter(
    client: TestClient, app_session_factory: sessionmaker[Session], reauthenticate: Any
) -> LoggedIn:
    who = sign_in(client, app_session_factory)
    reauthenticate(who)
    return who


@pytest.fixture
def lines(app_session_factory: sessionmaker[Session]) -> Lines:
    found = Lines(app_session_factory)
    found.add(
        NOW - timedelta(hours=1),
        "fg_web",
        srcip="172.20.41.105",
        hostname="=cmd|' /C calc'!A0",
        url="+1+1",
        action="blocked",
        catdesc="@SUM(1)",
        text="-2+3 \x01 kontrol",
    )
    found.add(
        NOW - timedelta(hours=2),
        "fg_web",
        srcip="172.20.41.105",
        hostname="www.example.com",
        action="passthrough",
        text="ok",
    )
    found.add(
        NOW - timedelta(hours=3),
        "fg_web",
        srcip="172.20.41.105",
        hostname="haber.example",
        action="passthrough",
        text="ok",
    )
    return found


def search_id(client: TestClient, q: str = "srcip:172.20.41.105") -> str:
    response = client.post(API, json={"q": q, "from": START.isoformat(), "to": END.isoformat()})
    assert response.status_code == 200, response.text
    return str(response.json()["search_id"])


def export(client: TestClient, search: str, fmt: str = "csv") -> Any:
    return client.post(f"{API}/export", json={"search_id": search, "format": fmt})


def audits(factory: sessionmaker[Session]) -> list[AuditLog]:
    with factory() as session:
        rows = list(
            session.scalars(
                select(AuditLog).where(AuditLog.action == "search.export").order_by(AuditLog.id)
            )
        )
        for row in rows:
            session.expunge(row)
    return rows


def _digest(response: Any) -> str:
    header = response.headers["Repr-Digest"]
    assert header.startswith("sha-256=:")
    return base64.b64decode(header.removeprefix("sha-256=:").removesuffix(":")).hex()


def test_csv(
    client: TestClient, exporter: LoggedIn, lines: Lines, app_session_factory: sessionmaker[Session]
) -> None:
    response = export(client, search_id(client))
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == "text/csv; charset=utf-8"
    name = response.headers["content-disposition"]
    assert name.startswith('attachment; filename="arama-')
    assert name.endswith('.csv"')
    assert "-kismi" not in name
    assert "X-Arci-Export-Truncated" not in response.headers
    data = response.content
    assert data.startswith("﻿".encode())
    assert int(response.headers["Content-Length"]) == len(data)
    sha256 = hashlib.sha256(data).hexdigest()
    assert _digest(response) == sha256
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"), newline=""), delimiter=";"))
    assert rows[0][:6] == ["Zaman", "Alınma zamanı", "Kaynak", "Gönderici", "Tür", "Alt tür"]
    assert rows[0][-2:] == ["Metin", "Ham kayıt"]
    assert len(rows) == 4
    first = dict(zip(rows[0], rows[1], strict=True))
    assert first["Alan adı"] == "'=cmd|' /C calc'!A0"  # neutralised
    assert first["URL"] == "'+1+1"
    assert first["Kategori"] == "'@SUM(1)"
    assert first["Metin"] == "'-2+3 \\u0001 kontrol"  # display-safe and neutralised
    assert first["Tür"] == "Web filtresi"
    assert first["Ham kayıt"] == f"fortigate/{lines.written[0].day.isoformat()}/1/0"
    assert b"\r\n" in data
    (entry,) = audits(app_session_factory)
    assert entry.details["sha256"] == sha256
    assert entry.details["bytes"] == len(data)
    assert entry.details["rows"] == 3
    assert entry.details["format"] == "csv"
    assert entry.details["truncated"] is None
    assert entry.details["q"] == "srcip:172.20.41.105"
    assert "www.example.com" not in str(entry.details)


def test_xlsx(
    client: TestClient, exporter: LoggedIn, lines: Lines, app_session_factory: sessionmaker[Session]
) -> None:
    response = export(client, search_id(client), "xlsx")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert _digest(response) == hashlib.sha256(response.content).hexdigest()
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Sonuçlar", "Bilgi"]
    results = book["Sonuçlar"]
    values = [[cell.value for cell in row] for row in results.iter_rows()]
    assert values[0][0] == "Zaman"
    assert len(values) == 4
    formula = next(cell for cell in results[2] if cell.value == "=cmd|' /C calc'!A0")
    assert formula.data_type == "s"  # text, never a formula
    assert "\\u0001" in str(results[2][-2].value)  # a control character written visibly
    info = {row[0].value: row[1].value for row in book["Bilgi"].iter_rows()}
    assert info["Sorgu / Query"] == "srcip:172.20.41.105"
    assert info["Satır / Rows"] == "3"
    assert info["Not"] == NOT_SEALED[0]
    assert info["Note"] == NOT_SEALED[1]
    assert info["Kullanıcı / User"] == exporter.username
    assert (
        audits(app_session_factory)[0].details["sha256"]
        == hashlib.sha256(response.content).hexdigest()
    )


def test_a_file_cut_at_the_row_limit_says_so(
    client: TestClient, exporter: LoggedIn, lines: Lines, app_session_factory: sessionmaker[Session]
) -> None:
    setting(app_session_factory, "search.export_max_rows", 2)
    response = export(client, search_id(client))
    assert response.headers["X-Arci-Export-Truncated"] == "rows"
    assert '-kismi.csv"' in response.headers["content-disposition"]
    assert len(response.content.decode("utf-8-sig").strip().split("\r\n")) == 3
    assert audits(app_session_factory)[0].details["truncated"] == "rows"
    book = load_workbook(io.BytesIO(export(client, search_id(client), "xlsx").content))
    info = {row[0].value: row[1].value for row in book["Bilgi"].iter_rows()}
    assert info["Satır / Rows"] == "2"
    assert info["Durum / Status"].startswith("Satır sınırında kesildi")


class FakeClock:
    def __init__(self, step: float) -> None:
        self.step = step
        self.now = 0.0

    def __call__(self) -> float:
        self.now += self.step
        return self.now


def test_a_file_cut_at_the_time_limit_says_so(
    app: FastAPI,
    client: TestClient,
    exporter: LoggedIn,
    lines: Lines,
    app_session_factory: sessionmaker[Session],
) -> None:
    setting(app_session_factory, "search.export_budget_seconds", 10)
    search = search_id(client)
    state_of(app).clock = FakeClock(4.0)  # the budget (10 s) is gone after a window
    response = export(client, search)
    assert response.status_code == 200, response.text
    assert response.headers["X-Arci-Export-Truncated"] == "time"
    assert '-kismi.csv"' in response.headers["content-disposition"]
    assert audits(app_session_factory)[0].details["truncated"] == "time"


def test_export_needs_a_recent_re_authentication(
    client: TestClient, app_session_factory: sessionmaker[Session], lines: Lines
) -> None:
    sign_in(client, app_session_factory)
    response = export(client, search_id(client))
    assert response.status_code == 403
    assert response.json()["error"]["code"] == "reauth_required"
    assert audits(app_session_factory) == []


def test_one_export_at_a_time(
    app: FastAPI, client: TestClient, exporter: LoggedIn, lines: Lines
) -> None:
    search = search_id(client)
    with state_of(app).exports.hold(exporter.user_id):
        response = export(client, search)
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "search_export_running"
    assert export(client, search).status_code == 200


def test_the_hourly_limit_survives_a_new_application(
    api_config: Config,
    db_engine: Engine,
    client: TestClient,
    exporter: LoggedIn,
    lines: Lines,
    app_session_factory: sessionmaker[Session],
) -> None:
    setting(app_session_factory, "search.export_max_per_hour", 2)
    search = search_id(client)
    assert export(client, search).status_code == 200
    assert export(client, search).status_code == 200
    refused = export(client, search)
    assert refused.status_code == 429
    error = refused.json()["error"]
    assert error["code"] == "search_export_rate_limited"
    assert "2" in error["message_tr"]
    assert int(refused.headers["Retry-After"]) > 3000
    # A restarted API counts the same audit entries.
    with TestClient(
        create_app(api_config, db_engine, audit_flush_seconds=None), base_url="http://testserver"
    ) as again:
        again.cookies.update(client.cookies)
        again.headers.update(client.headers)
        fresh = again.post(
            API, json={"q": "", "from": START.isoformat(), "to": END.isoformat()}
        ).json()["search_id"]
        assert export(again, fresh).json()["error"]["code"] == "search_export_rate_limited"


def test_another_session_cannot_export_the_search(
    app: FastAPI,
    client: TestClient,
    exporter: LoggedIn,
    lines: Lines,
    app_session_factory: sessionmaker[Session],
    reauthenticate: Any,
) -> None:
    search = search_id(client)
    with TestClient(app, base_url="http://testserver") as other:
        who = sign_in(other, app_session_factory)
        reauthenticate(who)
        response = export(other, search)
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "search_expired"
