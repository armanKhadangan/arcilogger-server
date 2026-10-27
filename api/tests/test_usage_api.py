"""``/api/v1/usage`` on the golden usage site (worker/tests/test_usage_golden.py has the numbers).

The expected values below are the hand-computed ones of the golden test (14.10.2026): company
active time 1020 s, online 1440 s, download 734500 bytes, 14 sessions, 3 blocked attempts; Ayşe
360 s active (240 in work hours), PC-02 300 s, Burak 180 s, Can 120 s, the bare address 60 s.
"""

from __future__ import annotations

import csv
import io
from collections.abc import Callable
from datetime import datetime, timedelta
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select, update
from sqlalchemy.orm import Session, sessionmaker

from arcilogger_api.auth.accounts import accept_privacy, privacy_version
from arcilogger_api.testing import LoggedIn
from arcilogger_core.config import Config
from arcilogger_core.layout import Layout
from arcilogger_core.secretstore import SecretStore
from arcilogger_core.timeutil import site_day, utc_now
from arcilogger_db.audit import canonical_entry, entry_hash
from arcilogger_db.models import AuditLog, Person, UsageRebuild
from arcilogger_db.settings import set_setting
from arcilogger_worker.testing import tr
from arcilogger_worker.usage import query
from arcilogger_worker.usage.jobs import process_requests, rebuild_day
from arcilogger_worker.usage.pseudonyms import pseudonym
from arcilogger_worker.usage.testing import (
    BARE,
    DAY,
    NEXT,
    PC2_MAC,
    GoldenSite,
    build_golden_site,
)

API = "/api/v1/usage"
PERIOD = {"from": DAY.isoformat(), "to": DAY.isoformat()}
NOW = tr(NEXT + timedelta(days=1), 9)
Login = Callable[..., LoggedIn]
PSEUDONYM_KEY = "golden-usage-pseudonym-key-0123456789"  # gitleaks:allow (test-only key)


@pytest.fixture
def site(db_session: Session, layout: Layout) -> GoldenSite:
    golden = build_golden_site(db_session, layout)
    rebuild_day(db_session, DAY, now=NOW)
    rebuild_day(db_session, NEXT, now=NOW)
    return golden


def _signed_in(
    client: TestClient,
    login_as: Login,
    factory: sessionmaker[Session],
    role: str,
    *,
    notice: bool = True,
) -> LoggedIn:
    who = login_as(client, role)
    if notice:
        with factory() as session:
            accept_privacy(session, who.user_id, privacy_version(session), ip=None)
            session.commit()
    return who


@pytest.fixture
def admin(
    client: TestClient,
    login_as: Login,
    app_session_factory: sessionmaker[Session],
    site: GoldenSite,
) -> TestClient:
    _signed_in(client, login_as, app_session_factory, "admin")
    return client


def _get(client: TestClient, path: str, **params: Any) -> Any:
    response = client.get(f"{API}{path}", params={**PERIOD, **params})
    assert response.status_code == 200, response.text
    return response.json()


def test_company_totals_match_the_golden_day(admin: TestClient) -> None:
    data = _get(admin, "/company")
    assert data["estimated"] is True
    assert data["notice_key"] == "usage.method_note"
    assert data["period"] == {"from": "2026-10-14", "to": "2026-10-14", "hours": "all"}
    assert data["computed_until"] == "2026-10-14T21:00:00Z"
    assert data["active_users"] == 5
    metrics = data["metrics"]
    assert {k: metrics[k] for k in metrics if not k.endswith("activity")} == {
        "online_seconds": 1440,
        "active_seconds": 1020,
        "download_bytes": 734500,
        "upload_bytes": 11560,
        "sessions": 14,
        "blocked": 3,
        "off_hours_active_seconds": 180,
    }
    assert metrics["first_activity"] == "2026-10-14T06:00:00Z"  # 09:00 site time
    assert len(data["hourly"]) == 24
    nine = data["hourly"][9]
    assert nine == {
        "hour": "2026-10-14T06:00:00Z",
        "active_users": 1,
        "active_seconds": 240,
        "download_bytes": 621600,
        "upload_bytes": 8160,
    }
    sites = {row["domain"]: row for row in data["top_sites"]}
    # video 210, PC-02's bulut 240, portal 180 + 120 …; background sites have no time.
    assert sites["bulut.example"]["seconds"] == 240
    assert sites["portal.example"]["seconds"] == 300
    assert sites["portal.example"]["users"] == 2
    assert sites["video.example"]["category"] == "Video"
    assert "windowsupdate.com" not in sites
    assert sum(row["seconds"] for row in data["top_sites"]) == 1020
    categories = {row["category"]: row for row in data["top_categories"]}
    assert categories["İş Uygulamaları"]["seconds"] == 300
    assert categories["Bilinmiyor"]["seconds"] == 90  # posta.example 30 + 203.0.113.110 60
    departments = {row["department"]: row for row in data["departments"]}
    assert departments["Finans"]["source"] == "person"
    # Ayşe 360 s and Can 120 s (persons) + the bare address of VLAN 41, 60 s.
    assert departments["Finans"]["metrics"]["active_seconds"] == 540
    assert departments["Finans"]["active_users"] == 3
    assert departments["İdari İşler"]["source"] == "vlan"  # PC-02: VLAN 43
    assert departments["Planlama"]["source"] == "person"
    apps = {row["app"]: row for row in data["top_apps"]}
    assert apps["SSL_TLSv1.3.PQC"] == {
        "app": "SSL_TLSv1.3.PQC",
        "appcat": "Network.Service",
        "seconds": 210,
        "bytes": 20460,
    }


def test_work_and_off_hours_split_the_day(admin: TestClient) -> None:
    work = _get(admin, "/company", hours="work")["metrics"]
    off = _get(admin, "/company", hours="off")["metrics"]
    assert (work["active_seconds"], off["active_seconds"]) == (840, 180)
    assert (work["online_seconds"], off["online_seconds"]) == (1260, 180)
    assert off["off_hours_active_seconds"] == 180
    assert work["off_hours_active_seconds"] == 0


def test_a_period_without_an_end_runs_to_today(
    admin: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    # The router reads the test's moment: a site midnight (21:00 UTC) between the test's `today`
    # and a request would otherwise move the default `to` (and the refused `from`) a day on.
    moment = utc_now()
    monkeypatch.setattr("arcilogger_api.routers.usage.utc_now", lambda: moment)
    today = site_day(moment)
    since = today - timedelta(days=3)
    period = admin.get(f"{API}/company", params={"from": since.isoformat()}).json()["period"]
    assert period == {"from": since.isoformat(), "to": today.isoformat(), "hours": "all"}
    period = admin.get(f"{API}/company", params={"to": DAY.isoformat()}).json()["period"]
    assert (period["from"], period["to"]) == (DAY.isoformat(), DAY.isoformat())
    period = admin.get(f"{API}/company").json()["period"]
    assert (period["from"], period["to"]) == (today.isoformat(), today.isoformat())
    future = admin.get(f"{API}/company", params={"from": (today + timedelta(days=1)).isoformat()})
    assert (future.status_code, future.json()["error"]["code"]) == (422, "usage_period")


def test_subject_lists(admin: TestClient, site: GoldenSite) -> None:
    persons = _get(admin, "/subjects", kind="person")
    assert persons["total"] == 3
    assert [row["subject"]["label"] for row in persons["items"]] == [
        "Ayşe Yılmaz",
        "Burak Demir",
        "Can Öztürk",
    ]
    ayse = persons["items"][0]
    assert ayse["subject"] == {
        "kind": "person",
        "key": f"person:{site.ayse}",
        "label": "Ayşe Yılmaz",
        "pseudonymous": False,
        "mac": None,
        "ip": "172.20.43.37",
        "hostname": None,
    }
    assert ayse["department"] == "Finans"
    assert ayse["metrics"]["active_seconds"] == 360
    assert ayse["top_sites"] == [
        {"domain": "video.example", "seconds": 210},
        {"domain": "haber.example", "seconds": 120},
        {"domain": "posta.example", "seconds": 30},
    ]
    assert ayse["top_categories"][0] == {"category": "Video", "seconds": 210}
    devices = _get(admin, "/subjects", kind="device", sort="label", order="asc")
    assert [row["subject"]["key"] for row in devices["items"]] == [
        "ip:172.20.41.77",
        f"device:{PC2_MAC}",
    ]
    assert devices["items"][1]["subject"]["label"] == "PC-02"
    paged = _get(admin, "/subjects", kind="person", limit=1, offset=1)
    assert (paged["total"], paged["limit"], paged["offset"]) == (3, 1, 1)
    assert [row["subject"]["label"] for row in paged["items"]] == ["Burak Demir"]
    found = _get(admin, "/subjects", kind="person", q="AYŞE")
    assert [row["subject"]["label"] for row in found["items"]] == ["Ayşe Yılmaz"]
    departments = _get(admin, "/subjects", kind="department")
    assert [row["department"] for row in departments["items"]][:2] == ["Finans", "İdari İşler"]
    # Department filter: Finans = Ayşe, Can and the bare address of VLAN 41.
    finance = _get(admin, "/subjects", kind="person", department="Finans")
    assert [row["subject"]["label"] for row in finance["items"]] == ["Ayşe Yılmaz", "Can Öztürk"]
    devices = _get(admin, "/subjects", kind="device", department="Finans")
    assert [row["subject"]["key"] for row in devices["items"]] == [f"ip:{BARE}"]
    sites = _get(admin, "/sites", department="Finans")
    # Ayşe's seven sites (four of them background traffic without time), Can's portal.example
    # and the bare address's destination.
    assert [(row["domain"], row["seconds"]) for row in sites["items"]] == [
        ("video.example", 210),
        ("haber.example", 120),
        ("portal.example", 120),
        ("203.0.113.110", 60),
        ("posta.example", 30),
        ("windowsupdate.com", 0),
        ("Microsoft-Azure", 0),
        ("8.8.8.8", 0),
        ("reklam.example", 0),
    ]


def test_subject_detail_of_a_person_day(admin: TestClient, site: GoldenSite) -> None:
    data = _get(admin, f"/subjects/person:{site.ayse}")
    assert data["metrics"]["active_seconds"] == 360
    assert data["department"] == "Finans"
    assert [device["label"] for device in data["devices"]] == ["Ayşe'nin bilgisayarı"]
    # 09:00-09:02 Video, 09:02 posta (no category), 09:03 Video, 19:30-19:32 Haber.
    assert [(t["start"], t["end"], t["category"]) for t in data["timeline"]] == [
        ("2026-10-14T06:00:00Z", "2026-10-14T06:02:00Z", "Video"),
        ("2026-10-14T06:02:00Z", "2026-10-14T06:03:00Z", None),
        ("2026-10-14T06:03:00Z", "2026-10-14T06:04:00Z", "Video"),
        ("2026-10-14T16:30:00Z", "2026-10-14T16:32:00Z", "Haber"),
    ]
    sites = {row["domain"]: row for row in data["sites"]}
    assert len(sites) == 7
    assert sites["video.example"]["share"] == round(210 / 360, 4)
    assert sites["windowsupdate.com"]["background"] is True
    assert sites["Microsoft-Azure"]["kind"] == "service"
    hosts = {row["hostname"]: row for row in data["hostnames"]}
    assert hosts["www.video.example"] == {
        "hostname": "www.video.example",
        "domain": "video.example",
        "seconds": 210,
        "bytes": 20460,
        "sessions": 1,
    }
    assert [row["url"] for row in data["urls"]] == ["/gundem", "/canli", "/izle?v=42"]
    assert data["urls"][2]["hits"] == 2
    assert [(b["source"], b["category"]) for b in data["blocked"]] == [
        ("app", "P2P"),
        ("dns", "Kumar"),
        ("web", "Kumar"),
    ]
    assert len(data["hourly"]) == 24
    assert data["hourly"][9]["top_domain"] == "video.example"
    (device,) = data["per_device"]
    assert device["metrics"]["active_seconds"] == 360
    assert device["metrics"]["blocked"] == 3
    assert device["top_sites"][0] == {"domain": "video.example", "seconds": 210}
    assert len(data["heatmap"]) == 30 * 24
    cell = next(c for c in data["heatmap"] if c["day"] == "2026-10-14" and c["hour"] == 9)
    assert cell["active_minutes"] == 4
    # Viewing a person's report is audited (spec 8.5), without the name.
    with Session(admin.app.state.arci.engine) as session:  # type: ignore[attr-defined]
        entry = session.scalars(
            select(AuditLog).where(AuditLog.action == "usage.view").order_by(AuditLog.id.desc())
        ).first()
    assert entry is not None
    assert entry.target == f"usage:person:{site.ayse}"
    assert "Ayşe" not in str(entry.details)


def test_reassigned_device_and_unknown_subjects(admin: TestClient, site: GoldenSite) -> None:
    burak = _get(admin, f"/subjects/person:{site.burak}")
    can = _get(admin, f"/subjects/person:{site.can}")
    assert burak["metrics"]["sessions"] == can["metrics"]["sessions"] == 2
    assert burak["per_device"][0]["device"]["mac"] == can["per_device"][0]["device"]["mac"]
    response = admin.get(f"{API}/subjects/person:999999", params=PERIOD)
    assert (response.status_code, response.json()["error"]["code"]) == (
        404,
        "usage_subject_not_found",
    )


def test_sites_and_categories(admin: TestClient, site: GoldenSite) -> None:
    sites = _get(admin, "/sites", sort="download_bytes")
    assert sites["items"][0]["domain"] == "windowsupdate.com"
    assert sites["total"] == 11  # 7 of Ayşe, 2 of PC-02, portal.example, 203.0.113.110
    video = _get(admin, "/sites", q="vid")["items"]
    assert [row["domain"] for row in video] == ["video.example"]
    detail = _get(admin, "/sites/portal.example")
    assert detail["totals"]["seconds"] == 300
    assert [(h["hostname"], h["seconds"]) for h in detail["hostnames"]] == [("portal.example", 300)]
    assert [(u["subject"]["label"], u["metrics"]["active_seconds"]) for u in detail["users"]] == [
        ("Burak Demir", 180),
        ("Can Öztürk", 120),
    ]
    categories = _get(admin, "/categories")
    assert {row["category"] for row in categories["items"]} >= {"Video", "Bilinmiyor"}
    video_category = _get(admin, "/categories/Video")
    assert video_category["totals"]["seconds"] == 210
    assert [row["domain"] for row in video_category["sites"]] == ["video.example"]
    assert video_category["users"][0]["subject"]["label"] == "Ayşe Yılmaz"
    unknown = _get(admin, "/categories/Bilinmiyor")
    assert unknown["totals"]["seconds"] == 90
    assert admin.get(f"{API}/sites/none.example", params=PERIOD).status_code == 404


def test_excluded_persons_leave_person_level_views(admin: TestClient, site: GoldenSite) -> None:
    # Set where the administrator sets it (Cihazlar ve Kişiler), read when a view is read: no
    # rebuild, also for the days already rolled up.
    excluded = admin.patch(f"/api/v1/persons/{site.ayse}", json={"usage_excluded": True})
    assert excluded.status_code == 200, excluded.text
    persons = _get(admin, "/subjects", kind="person")
    assert [row["subject"]["label"] for row in persons["items"]] == ["Burak Demir", "Can Öztürk"]
    assert admin.get(f"{API}/subjects/person:{site.ayse}", params=PERIOD).status_code == 404
    video = _get(admin, "/sites/video.example")
    assert video["users"] == []
    assert video["totals"]["seconds"] == 210  # the company totals still count the traffic
    assert _get(admin, "/company")["metrics"]["active_seconds"] == 1020
    included = admin.patch(f"/api/v1/persons/{site.ayse}", json={"usage_excluded": False})
    assert included.status_code == 200, included.text
    persons = _get(admin, "/subjects", kind="person")
    assert persons["items"][0]["subject"]["label"] == "Ayşe Yılmaz"


def test_exports_are_audited_files(admin: TestClient, site: GoldenSite) -> None:
    response = admin.get(f"{API}/export.csv", params={**PERIOD, "view": "subjects"})
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    text = response.content.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text), delimiter=";"))
    assert rows[0][:3] == ["Ad", "Tür", "Anahtar"]
    assert [row[0] for row in rows[1:]] == ["Ayşe Yılmaz", "Burak Demir", "Can Öztürk"]
    response = admin.get(
        f"{API}/export.xlsx", params={**PERIOD, "view": "subject", "key": f"person:{site.ayse}"}
    )
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames[:3] == ["Siteler", "Özet", "Kategoriler"]
    assert "Yöntem" in book.sheetnames
    first = [cell.value for cell in book["Siteler"][2]]
    assert first[:4] == ["video.example", "domain", "Video", 210]
    company = admin.get(f"{API}/export.csv", params={**PERIOD, "view": "company"})
    lines = company.content.decode("utf-8-sig").splitlines()
    assert lines[1].startswith("Şirket (toplam)")
    missing = admin.get(f"{API}/export.csv", params={**PERIOD, "view": "site"})
    assert missing.status_code == 422
    with Session(admin.app.state.arci.engine) as session:  # type: ignore[attr-defined]
        entries = list(
            session.scalars(
                select(AuditLog).where(AuditLog.action == "usage.export").order_by(AuditLog.id)
            )
        )
    assert [(e.details["view"], e.details["format"]) for e in entries] == [
        ("subjects", "csv"),
        ("subject", "xlsx"),
        ("company", "csv"),
    ]
    assert entries[0].details["rows"] == 3


def test_site_and_category_users_are_exported_whole_with_their_own_figures(
    admin: TestClient, site: GoldenSite, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(query, "USERS_MAX", 1)
    detail = _get(admin, "/sites/portal.example")
    assert [u["subject"]["label"] for u in detail["users"]] == ["Burak Demir"]
    assert detail["users_total"] == 2
    response = admin.get(
        f"{API}/export.csv", params={**PERIOD, "view": "site", "domain": "portal.example"}
    )
    assert response.status_code == 200
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig")), delimiter=";"))
    assert rows[0] == [
        "Ad", "Tür", "Anahtar", "MAC", "IP", "Departman", "Süre (sn, tahmini)", "Ham süre (sn)",
        "Mesai dışı süre (sn, tahmini)", "İndirilen (bayt)", "Yüklenen (bayt)", "Oturum",
    ]  # fmt: skip
    assert [(row[0], row[6], row[7], row[9], row[11]) for row in rows[1:]] == [
        ("Burak Demir", "180", "180", "16000", "2"),
        ("Can Öztürk", "120", "120", "7000", "2"),
    ]
    # A category: no online time (its sites' raw times would count a minute once per site).
    unknown = _get(admin, "/categories/Bilinmiyor")
    assert (unknown["users_total"], unknown["sites_total"]) == (2, 6)
    response = admin.get(
        f"{API}/export.xlsx", params={**PERIOD, "view": "category", "category": "Bilinmiyor"}
    )
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.content))["Kullanıcılar"]
    header, *users = ([cell.value for cell in row] for row in sheet.iter_rows())
    assert header == [
        "Ad", "Tür", "Anahtar", "MAC", "IP", "Departman", "Süre (sn, tahmini)",
        "Mesai dışı süre (sn, tahmini)", "İndirilen (bayt)", "Yüklenen (bayt)", "Oturum",
        "En çok kullanılan siteler",
    ]  # fmt: skip
    assert [row[0] for row in users] == ["172.20.41.77", "Ayşe Yılmaz"]
    assert users[1][6:] == [30, 0, 658900, 7700, 5, "posta.example (30 sn)"]


def test_whole_lists_come_in_one_answer(admin: TestClient, site: GoldenSite) -> None:
    """The web reads every site and subject at once: no page-by-page recomputation."""
    sites = _get(admin, "/sites", limit=100_000)
    assert (sites["total"], len(sites["items"]), sites["limit"]) == (11, 11, 100_000)
    persons = _get(admin, "/subjects", kind="person", limit=100_000)
    assert len(persons["items"]) == persons["total"] == 3
    too_many = admin.get(f"{API}/sites", params={**PERIOD, "limit": 100_001})
    assert too_many.status_code == 422


def test_formula_injection_is_neutralised(
    admin: TestClient, site: GoldenSite, app_session_factory: sessionmaker[Session]
) -> None:
    with app_session_factory() as session:
        session.execute(update(Person).where(Person.id == site.can).values(name="=HYPERLINK(1)"))
        # A control character cannot be stored in an Excel file: written visibly instead.
        session.execute(update(Person).where(Person.id == site.burak).values(name="Burak\x01"))
        session.commit()
    csv_text = admin.get(f"{API}/export.csv", params={**PERIOD, "view": "subjects"}).content
    assert b"'=HYPERLINK(1)" in csv_text
    book = load_workbook(
        io.BytesIO(admin.get(f"{API}/export.xlsx", params={**PERIOD, "view": "subjects"}).content)
    )
    values = [row[0].value for row in book.active.iter_rows(min_row=2)]
    cell = next(c for row in book.active.iter_rows(min_row=2) for c in row[:1] if c.value)
    assert "=HYPERLINK(1)" in values
    assert "Burak\\x01" in values
    assert cell.data_type == "s"


def test_rebuild_requests_run_in_the_worker(
    admin: TestClient, app_session_factory: sessionmaker[Session]
) -> None:
    response = admin.post(f"{API}/rebuild", json={"from": DAY.isoformat(), "to": NEXT.isoformat()})
    assert response.status_code == 202, response.text
    job = response.json()
    assert (job["status"], job["progress"], job["from"], job["to"]) == (
        "queued",
        0.0,
        "2026-10-14",
        "2026-10-15",
    )
    report = process_requests(app_session_factory, now=NOW)
    assert len(report.days) == 2
    status = admin.get(f"{API}/rebuild/{job['job_id']}").json()
    assert (status["status"], status["progress"]) == ("done", 1.0)
    bad = admin.post(f"{API}/rebuild", json={"from": NEXT.isoformat(), "to": DAY.isoformat()})
    assert (bad.status_code, bad.json()["error"]["code"]) == (422, "usage_rebuild_range")
    assert admin.get(f"{API}/rebuild/99999").status_code == 404
    with app_session_factory() as session:
        assert session.scalars(select(UsageRebuild)).one().status == "done"


def test_no_rebuild_before_the_oldest_log(
    admin: TestClient, app_session_factory: sessionmaker[Session]
) -> None:
    """Without its logs a day would be rebuilt empty: nothing is queued, its rollups stay."""
    early = admin.post(
        f"{API}/rebuild",
        json={"from": (DAY - timedelta(days=30)).isoformat(), "to": DAY.isoformat()},
    )
    assert (early.status_code, early.json()["error"]["code"]) == (422, "usage_rebuild_no_logs")
    assert "14.10.2026" in early.json()["error"]["message_tr"]
    with app_session_factory() as session:
        assert session.scalars(select(UsageRebuild)).first() is None
    # The day of the oldest log itself can be rebuilt.
    first = admin.post(f"{API}/rebuild", json={"from": DAY.isoformat(), "to": DAY.isoformat()})
    assert first.status_code == 202, first.text


def test_without_any_log_only_today_can_be_rebuilt(
    client: TestClient,
    login_as: Login,
    app_session_factory: sessionmaker[Session],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """No ``fg_traffic`` row at all: every day before today is refused, today is not."""
    _signed_in(client, login_as, app_session_factory, "admin")
    # The router reads the test's moment: a site midnight (21:00 UTC) between here and a
    # request would otherwise put the router's "today" (and its first rebuildable day) one day
    # after the test's.
    moment = utc_now()
    monkeypatch.setattr("arcilogger_api.routers.usage.utc_now", lambda: moment)
    today = site_day(moment)
    yesterday = (today - timedelta(days=1)).isoformat()
    early = client.post(f"{API}/rebuild", json={"from": yesterday, "to": yesterday})
    assert (early.status_code, early.json()["error"]["code"]) == (422, "usage_rebuild_no_logs")
    assert f"{today:%d.%m.%Y}" in early.json()["error"]["message_tr"]
    now = client.post(f"{API}/rebuild", json={"from": today.isoformat(), "to": today.isoformat()})
    assert now.status_code == 202, now.text


def test_roles_and_the_privacy_notice(
    app: Any,
    login_as: Login,
    app_session_factory: sessionmaker[Session],
    site: GoldenSite,
) -> None:
    with TestClient(app, base_url="http://testserver") as hr:
        _signed_in(hr, login_as, app_session_factory, "hr", notice=False)
        refused = hr.get(f"{API}/subjects", params=PERIOD)
        assert (refused.status_code, refused.json()["error"]["code"]) == (
            403,
            "privacy_notice_required",
        )
        assert hr.get(f"{API}/company", params=PERIOD).status_code == 200
    with TestClient(app, base_url="http://testserver") as viewer:
        _signed_in(viewer, login_as, app_session_factory, "viewer")
        # Owner decision 2026-10-20: names for every role that may see usage.
        names = viewer.get(f"{API}/subjects", params=PERIOD).json()
        assert names["items"][0]["subject"]["label"] == "Ayşe Yılmaz"
        assert (
            viewer.post(
                f"{API}/rebuild", json={"from": "2026-10-14", "to": "2026-10-14"}
            ).status_code
            == 403
        )
        assert viewer.get(f"{API}/pseudonyms").status_code == 403


def test_the_viewer_accepts_the_privacy_notice_too(
    app: Any,
    login_as: Login,
    app_session_factory: sessionmaker[Session],
    site: GoldenSite,
) -> None:
    """The viewer sees names (owner decision 2026-10-20), so it accepts the notice like the rest."""
    with TestClient(app, base_url="http://testserver") as viewer:
        _signed_in(viewer, login_as, app_session_factory, "viewer", notice=False)
        me = viewer.get("/api/v1/auth/me").json()
        assert (me["next"], me["privacy_notice"]["required"]) == ("privacy_notice", True)
        for path, params in (
            ("/subjects", {}),
            ("/subjects", {"kind": "device"}),
            (f"/subjects/person:{site.ayse}", {}),
            ("/sites", {}),
            ("/sites/portal.example", {}),
            ("/categories", {}),
            ("/categories/Video", {}),
            ("/export.csv", {"view": "subjects"}),
            ("/export.xlsx", {"view": "subject", "key": f"person:{site.ayse}"}),
        ):
            refused = viewer.get(f"{API}{path}", params={**PERIOD, **params})
            assert (refused.status_code, refused.json()["error"]["code"]) == (
                403,
                "privacy_notice_required",
            ), (path, params)
        # The company totals name nobody: no notice needed for them.
        assert viewer.get(f"{API}/company", params=PERIOD).status_code == 200
        company = viewer.get(f"{API}/export.csv", params={**PERIOD, "view": "company"})
        assert company.status_code == 200


def test_optional_pseudonyms_for_the_anonymous_role(
    app: Any,
    api_config: Config,
    login_as: Login,
    app_session_factory: sessionmaker[Session],
    site: GoldenSite,
    admin: TestClient,
) -> None:
    with app_session_factory() as session:
        set_setting(session, "usage.pseudonymize_anonymous_role", True, updated_by="test")
        session.commit()
    with TestClient(app, base_url="http://testserver") as viewer:
        _signed_in(viewer, login_as, app_session_factory, "viewer")
        missing = viewer.get(f"{API}/subjects", params=PERIOD)
        assert (missing.status_code, missing.json()["error"]["code"]) == (
            409,
            "pseudonym_key_missing",
        )
        store = SecretStore(
            api_config.secrets_dir / "secret.key", api_config.secrets_dir / "secrets.json"
        )
        if not (api_config.secrets_dir / "secret.key").exists():
            SecretStore.create_key(api_config.secrets_dir / "secret.key")
        store.set("usage.pseudonym_key", PSEUDONYM_KEY)
        page = viewer.get(f"{API}/subjects", params=PERIOD).json()
        first = page["items"][0]["subject"]
        token, label = pseudonym(PSEUDONYM_KEY.encode(), f"person:{site.ayse}")
        assert first == {
            "kind": "person",
            "key": token,
            "label": label,
            "pseudonymous": True,
            "mac": None,
            "ip": None,
            "hostname": None,
        }
        detail = viewer.get(f"{API}/subjects/{token}", params=PERIOD)
        assert detail.status_code == 200
        assert detail.json()["subject"]["label"] == label
        assert viewer.get(f"{API}/subjects/person:{site.ayse}", params=PERIOD).status_code == 404
        export = viewer.get(f"{API}/export.csv", params={**PERIOD, "view": "subject", "key": token})
        assert export.status_code == 200
    # The audit log names the real subject (its id, never the name): "who viewed this person"
    # finds the views and exports of pseudonymous viewers too, also after a key change.
    with app_session_factory() as session:
        entries = session.scalars(
            select(AuditLog)
            .where(AuditLog.action.in_(("usage.view", "usage.export")))
            .order_by(AuditLog.id)
        ).all()
        assert [(e.action, e.target, e.details["pseudonymous"]) for e in entries] == [
            ("usage.view", f"usage:person:{site.ayse}", True),
            ("usage.export", f"usage:person:{site.ayse}", True),
        ]
        assert not any(token in str(e.details) or label in str(e.details) for e in entries)
    # The admin still sees names, and the mapping (audited).
    assert _get(admin, "/subjects")["items"][0]["subject"]["label"] == "Ayşe Yılmaz"
    mapping = admin.get(f"{API}/pseudonyms", params={"q": "ayşe"}).json()
    assert mapping == [
        {
            "pseudonym": label,
            "kind": "person",
            "subject": {
                "kind": "person",
                "key": f"person:{site.ayse}",
                "label": "Ayşe Yılmaz",
                "pseudonymous": False,
                "mac": None,
                "ip": "172.20.43.37",
                "hostname": None,
            },
        }
    ]
    with app_session_factory() as session:
        session.execute(update(Person).where(Person.id == site.burak).values(usage_excluded=True))
        session.commit()
    # An excluded person never appears, so the mapping has nothing to explain for him.
    assert admin.get(f"{API}/pseudonyms", params={"q": "burak"}).json() == []
    assert len(admin.get(f"{API}/pseudonyms", params={"q": "can"}).json()) == 1


def _usage_entries(client: TestClient, **params: Any) -> list[dict[str, Any]]:
    response = client.get("/api/v1/audit", params={"action": "usage.", **params})
    assert response.status_code == 200, response.text
    items: list[dict[str, Any]] = response.json()["items"]
    return sorted(items, key=lambda item: item["id"])


def _recovered(entry: dict[str, Any], candidates: list[str]) -> str | None:
    """The candidate subject that gives the entry's ``hash`` from its ``prev_hash`` — what an
    answer with hashes lets a reader find without the pseudonym key (person ids are small
    numbers, search lists the MACs and IPs)."""
    if entry["hash"] is None or entry["prev_hash"] is None:
        return None
    for candidate in candidates:
        canonical = canonical_entry(
            entry_id=entry["id"],
            time=datetime.fromisoformat(entry["time"]),
            actor_id=entry["actor_id"],
            actor=entry["actor"],
            action=entry["action"],
            target=f"usage:{candidate}",
            details=entry["details"],
            ip=entry["ip"],
        )
        if entry_hash(entry["prev_hash"], canonical) == entry["hash"]:
            return candidate
    return None


def test_a_pseudonymous_auditor_reads_pseudonyms_in_the_audit_log(
    app: Any,
    api_config: Config,
    login_as: Login,
    app_session_factory: sessionmaker[Session],
    site: GoldenSite,
    admin: TestClient,
) -> None:
    """The auditor sees pseudonyms on the Kullanım pages and holds ``audit.read``: its own
    ``usage.view`` / ``usage.export`` entries store the real subject (MAC, IP, person id), and
    ``GET /audit`` shows it the pseudonym token instead. The stored entries stay as they are."""
    with app_session_factory() as session:
        set_setting(session, "usage.pseudonymize_anonymous_role", True, updated_by="test")
        session.commit()
    if not (api_config.secrets_dir / "secret.key").exists():
        SecretStore.create_key(api_config.secrets_dir / "secret.key")
    store = SecretStore(
        api_config.secrets_dir / "secret.key", api_config.secrets_dir / "secrets.json"
    )
    store.set("usage.pseudonym_key", PSEUDONYM_KEY)
    key = PSEUDONYM_KEY.encode()
    ayse, device, bare = f"person:{site.ayse}", f"device:{PC2_MAC}", f"ip:{BARE}"
    tokens = {real: pseudonym(key, real)[0] for real in (ayse, device, bare)}
    with TestClient(app, base_url="http://testserver") as auditor:
        _signed_in(auditor, login_as, app_session_factory, "auditor")
        devices = auditor.get(f"{API}/subjects", params={**PERIOD, "kind": "device"}).json()
        assert {tokens[device], tokens[bare]} <= {i["subject"]["key"] for i in devices["items"]}
        for real in (ayse, device, bare):
            assert auditor.get(f"{API}/subjects/{tokens[real]}", params=PERIOD).status_code == 200
        export = auditor.get(
            f"{API}/export.csv", params={**PERIOD, "view": "subject", "key": tokens[device]}
        )
        assert export.status_code == 200
        entries = _usage_entries(auditor)
        assert [(e["action"], e["target"], e["target_masked"]) for e in entries] == [
            ("usage.view", f"usage:{tokens[ayse]}", True),
            ("usage.view", f"usage:{tokens[device]}", True),
            ("usage.view", f"usage:{tokens[bare]}", True),
            ("usage.export", f"usage:{tokens[device]}", True),
        ]
        shown = str(entries)
        for real in (ayse, device, bare, PC2_MAC, BARE):
            assert real not in shown
        # No hash to test a guess against: every later hash depends on a masked entry, and two
        # chain heads around one view would do as well, so this reader gets none of them.
        candidates = [f"person:{n}" for n in range(1, site.ayse + 100)] + [device, bare]
        assert [_recovered(e, candidates) for e in entries] == [None] * 4
        everything = auditor.get("/api/v1/audit").json()["items"]
        assert {(e["prev_hash"], e["hash"]) for e in everything} == {(None, None)}
        status = auditor.get("/api/v1/audit/verify").json()
        assert (status["ok"], status["last_hash"]) == (True, None)
        # The target filter compares with what this reader sees: a real key finds nothing.
        found = _usage_entries(auditor, target=f"usage:{tokens[device]}")
        assert [e["action"] for e in found] == ["usage.view", "usage.export"]
        assert _usage_entries(auditor, target=f"usage:{device}") == []
        assert _usage_entries(auditor, target=f"usage:{ayse}") == []
        # Without the pseudonym key: a fixed placeholder, never the stored key.
        store.delete("usage.pseudonym_key")
        assert {e["target"] for e in _usage_entries(auditor)} == {"usage:masked"}
        assert {e["hash"] for e in _usage_entries(auditor)} == {None}
        assert len(_usage_entries(auditor, target="usage:masked")) == 4
        store.set("usage.pseudonym_key", PSEUDONYM_KEY)
        # The administrator sees names, so the stored keys ("who viewed this device").
        stored = _usage_entries(admin)
        assert [(e["target"], e["target_masked"]) for e in stored] == [
            (f"usage:{ayse}", False),
            (f"usage:{device}", False),
            (f"usage:{bare}", False),
            (f"usage:{device}", False),
        ]
        assert len(_usage_entries(admin, target=f"usage:{device}")) == 2
        # Its answer carries the hashes, from which the same search finds each subject.
        assert [_recovered(e, candidates) for e in stored] == [ayse, device, bare, device]
        newest = admin.get("/api/v1/audit").json()["items"][0]
        assert admin.get("/api/v1/audit/verify").json()["last_hash"] == newest["hash"]
        # With the setting off the auditor sees names again, and so the stored keys and hashes.
        with app_session_factory() as session:
            set_setting(session, "usage.pseudonymize_anonymous_role", False, updated_by="test")
            session.commit()
        assert _usage_entries(auditor) == stored
    with app_session_factory() as session:
        rows = session.scalars(
            select(AuditLog.target)
            .where(AuditLog.action.in_(("usage.view", "usage.export")))
            .order_by(AuditLog.id)
        ).all()
        assert list(rows) == [e["target"] for e in stored]
