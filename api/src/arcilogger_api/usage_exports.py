"""Usage exports: the rows of a ``/usage`` view as CSV or Excel (``GET /usage/export.*``).

* **CSV** (one table, Turkish Excel: ``;``, UTF-8 with BOM, cells neutralised against formula
  injection — :mod:`arcilogger_api.csvfiles`): the main table of the view (company: the
  departments with a company total first; one subject: its sites; one site: its users; one
  category: its sites).
* **Excel** (openpyxl, one sheet per section of the view, plus the sheet "Yöntem" with the method
  note: the numbers are estimates). Text cells are always written as text, never as formulas
  (host names and URLs come from the network) and control characters, which an Excel file
  cannot hold, are written visibly as ``\\xNN``; times in site time, durations in seconds and
  volumes in bytes as numbers, so they can be summed.

Lists are exported whole (no paging, none of the detail answers' caps); the view's filters
apply. The users of a site or category carry only the figures kept per site — time, a site's
raw time, off-hours time, volumes, sessions, a category's top sites —, never an online time of
a category, blocked attempts or first/last activity.
"""

from __future__ import annotations

import io
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Final

from fastapi import Response
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font

from arcilogger_api.csvfiles import csv_response, write_csv
from arcilogger_core.timeutil import to_site

XLSX_TYPE: Final = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
#: The method note (web locale ``usage.method_note``), Turkish and English.
METHOD_NOTE: Final = (
    "Tahmini değer: bağlantı kayıtlarından hesaplanır; arka planda açık sekmeler de sayılabilir.",
    "Estimated value: computed from connection logs; tabs left open in the background may be "
    "counted too.",
)
#: Excel's sheet name limit.
SHEET_NAME_MAX: Final = 31


@dataclass
class Sheet:
    name: str
    header: list[str]
    rows: list[list[Any]] = field(default_factory=list)


# -- columns --------------------------------------------------------------------------------------

METRIC_HEADER: Final = [
    "Çevrimiçi süre (sn)",
    "Aktif süre (sn, tahmini)",
    "Mesai dışı aktif (sn)",
    "İndirilen (bayt)",
    "Yüklenen (bayt)",
    "Oturum",
    "Engellenen",
    "İlk etkinlik",
    "Son etkinlik",
]
SITE_HEADER: Final = [
    "Alan adı",
    "Tür",
    "Kategori",
    "Süre (sn, tahmini)",
    "Ham süre (sn)",
    "Pay",
    "İndirilen (bayt)",
    "Yüklenen (bayt)",
    "Oturum",
    "Kullanıcı",
    "Arka plan",
]
CATEGORY_HEADER: Final = [
    "Kategori",
    "Süre (sn, tahmini)",
    "Pay",
    "İndirilen (bayt)",
    "Yüklenen (bayt)",
    "Kullanıcı",
]
SUBJECT_HEADER: Final = [
    "Ad",
    "Tür",
    "Anahtar",
    "MAC",
    "IP",
    "Departman",
    *METRIC_HEADER,
    "En çok kullanılan siteler",
    "En çok kullanılan kategoriler",
]
DEPARTMENT_HEADER: Final = ["Departman", "Kaynak", "Aktif kullanıcı", *METRIC_HEADER]
#: The users of one site or category: only the figures kept per site (no online time of a
#: category, no blocked attempts, no first/last activity).
USER_HEADER: Final = [
    "Ad",
    "Tür",
    "Anahtar",
    "MAC",
    "IP",
    "Departman",
    "Süre (sn, tahmini)",
    "Mesai dışı süre (sn, tahmini)",
    "İndirilen (bayt)",
    "Yüklenen (bayt)",
    "Oturum",
]
SITE_USER_HEADER: Final = [*USER_HEADER[:7], "Ham süre (sn)", *USER_HEADER[7:]]
CATEGORY_USER_HEADER: Final = [*USER_HEADER, "En çok kullanılan siteler"]


def metric_cells(metrics: Mapping[str, Any]) -> list[Any]:
    return [
        metrics["online_seconds"],
        metrics["active_seconds"],
        metrics["off_hours_active_seconds"],
        metrics["download_bytes"],
        metrics["upload_bytes"],
        metrics["sessions"],
        metrics["blocked"],
        metrics["first_activity"],
        metrics["last_activity"],
    ]


def site_cells(row: Mapping[str, Any]) -> list[Any]:
    return [
        row["domain"],
        row["kind"],
        row["category"],
        row["seconds"],
        row["raw_seconds"],
        row["share"],
        row["download_bytes"],
        row["upload_bytes"],
        row["sessions"],
        row["users"],
        row["background"],
    ]


def category_cells(row: Mapping[str, Any]) -> list[Any]:
    return [
        row["category"],
        row["seconds"],
        row["share"],
        row["download_bytes"],
        row["upload_bytes"],
        row["users"],
    ]


def subject_cells(row: Mapping[str, Any]) -> list[Any]:
    subject = row["subject"]
    return [
        subject["label"],
        subject["kind"],
        subject["key"],
        subject.get("mac"),
        subject.get("ip"),
        row["department"],
        *metric_cells(row["metrics"]),
        ", ".join(f"{s['domain']} ({s['seconds']} sn)" for s in row["top_sites"]),
        ", ".join(f"{c['category']} ({c['seconds']} sn)" for c in row["top_categories"]),
    ]


def user_cells(row: Mapping[str, Any], *, site: bool) -> list[Any]:
    """A row of :data:`SITE_USER_HEADER` (``site``) or :data:`CATEGORY_USER_HEADER`."""
    subject, metrics = row["subject"], row["metrics"]
    # A site user's online time is the site's raw time; a category has none.
    times = [metrics["active_seconds"], *([metrics["online_seconds"]] if site else [])]
    cells = [
        subject["label"],
        subject["kind"],
        subject["key"],
        subject.get("mac"),
        subject.get("ip"),
        row["department"],
        *times,
        metrics["off_hours_active_seconds"],
        metrics["download_bytes"],
        metrics["upload_bytes"],
        metrics["sessions"],
    ]
    if not site:
        cells.append(", ".join(f"{s['domain']} ({s['seconds']} sn)" for s in row["top_sites"]))
    return cells


def department_cells(row: Mapping[str, Any]) -> list[Any]:
    return [row["department"], row["source"], row["active_users"], *metric_cells(row["metrics"])]


def period_rows(data: Mapping[str, Any]) -> list[list[Any]]:
    period = data["period"]
    return [
        ["Başlangıç", period["from"]],
        ["Bitiş", period["to"]],
        ["Saatler", {"all": "tümü", "work": "mesai", "off": "mesai dışı"}[period["hours"]]],
        ["Hesaplanan son an", data["computed_until"]],
    ]


# -- views ----------------------------------------------------------------------------------------


def company_sheets(data: Mapping[str, Any]) -> list[Sheet]:
    summary = Sheet("Özet", ["Alan", "Değer"], period_rows(data))
    summary.rows.append(["Aktif kullanıcı", data["active_users"]])
    summary.rows += [
        list(pair) for pair in zip(METRIC_HEADER, metric_cells(data["metrics"]), strict=True)
    ]
    departments = Sheet("Departmanlar", DEPARTMENT_HEADER)
    departments.rows.append(
        ["Şirket (toplam)", "", data["active_users"], *metric_cells(data["metrics"])]
    )
    departments.rows += [department_cells(row) for row in data["departments"]]
    return [
        departments,
        summary,
        Sheet(
            "Saatlik",
            ["Saat", "Aktif kullanıcı", "Aktif süre (sn)", "İndirilen (bayt)", "Yüklenen (bayt)"],
            [
                [
                    p["hour"],
                    p["active_users"],
                    p["active_seconds"],
                    p["download_bytes"],
                    p["upload_bytes"],
                ]
                for p in data["hourly"]
            ],
        ),
        Sheet("Siteler", SITE_HEADER, [site_cells(row) for row in data["top_sites"]]),
        Sheet("Kategoriler", CATEGORY_HEADER, [category_cells(r) for r in data["top_categories"]]),
        Sheet(
            "Uygulamalar",
            ["Uygulama", "Uygulama kategorisi", "Süre (sn)", "Bayt"],
            [[a["app"], a["appcat"], a["seconds"], a["bytes"]] for a in data["top_apps"]],
        ),
    ]


def subjects_sheets(rows: Sequence[Mapping[str, Any]], *, departments: bool) -> list[Sheet]:
    if departments:
        return [Sheet("Departmanlar", DEPARTMENT_HEADER, [department_cells(r) for r in rows])]
    return [Sheet("Kişiler ve cihazlar", SUBJECT_HEADER, [subject_cells(r) for r in rows])]


def subject_sheets(data: Mapping[str, Any]) -> list[Sheet]:
    subject = data["subject"]
    summary = Sheet("Özet", ["Alan", "Değer"], period_rows(data))
    summary.rows += [
        ["Ad", subject["label"]],
        ["Tür", subject["kind"]],
        ["Anahtar", subject["key"]],
        ["MAC", subject.get("mac")],
        ["IP", subject.get("ip")],
        ["Departman", data["department"]],
        ["Cihazlar", ", ".join(d["label"] for d in data["devices"])],
    ]
    summary.rows += [
        list(pair) for pair in zip(METRIC_HEADER, metric_cells(data["metrics"]), strict=True)
    ]
    return [
        Sheet("Siteler", SITE_HEADER, [site_cells(row) for row in data["sites"]]),
        summary,
        Sheet("Kategoriler", CATEGORY_HEADER, [category_cells(r) for r in data["categories"]]),
        Sheet(
            "Uygulamalar",
            ["Uygulama", "Uygulama kategorisi", "Süre (sn)", "Bayt"],
            [[a["app"], a["appcat"], a["seconds"], a["bytes"]] for a in data["apps"]],
        ),
        Sheet(
            "Ana bilgisayarlar",
            ["Ana bilgisayar", "Alan adı", "Süre (sn)", "Bayt", "Oturum"],
            [
                [h["hostname"], h["domain"], h["seconds"], h["bytes"], h["sessions"]]
                for h in data["hostnames"]
            ],
        ),
        Sheet(
            "URL'ler",
            ["URL", "Ana bilgisayar", "İstek", "İlk", "Son", "Kategori", "Engellendi"],
            [
                [
                    u["url"],
                    u["hostname"],
                    u["hits"],
                    u["first_seen"],
                    u["last_seen"],
                    u["category"],
                    u["blocked"],
                ]
                for u in data["urls"]
            ],
        ),
        Sheet(
            "Engellenenler",
            [
                "Zaman",
                "Kaynak",
                "İşlem",
                "Alan adı",
                "Ana bilgisayar",
                "Kategori",
                "URL",
                "Uygulama",
            ],
            [
                [
                    b["time"],
                    b["source"],
                    b["action"],
                    b["domain"],
                    b.get("hostname"),
                    b["category"],
                    b.get("url"),
                    b.get("app"),
                ]
                for b in data["blocked"]
            ],
        ),
        Sheet(
            "Saatlik",
            [
                "Saat",
                "Aktif süre (sn)",
                "Çevrimiçi süre (sn)",
                "İndirilen (bayt)",
                "Yüklenen (bayt)",
                "En çok kullanılan site",
            ],
            [
                [
                    h["hour"],
                    h["active_seconds"],
                    h["online_seconds"],
                    h["download_bytes"],
                    h["upload_bytes"],
                    h["top_domain"],
                ]
                for h in data["hourly"]
            ],
        ),
        Sheet(
            "Cihazlar",
            ["Cihaz", "MAC", "IP", *METRIC_HEADER, "En çok kullanılan siteler"],
            [
                [
                    d["device"]["label"],
                    d["device"].get("mac"),
                    d["device"].get("ip"),
                    *metric_cells(d["metrics"]),
                    ", ".join(f"{s['domain']} ({s['seconds']} sn)" for s in d["top_sites"]),
                ]
                for d in data["per_device"]
            ],
        ),
        Sheet(
            "Zaman çizelgesi",
            ["Başlangıç", "Bitiş", "Kategori"],
            [[t["start"], t["end"], t["category"]] for t in data["timeline"]],
        ),
    ]


def sites_sheets(rows: Sequence[Mapping[str, Any]]) -> list[Sheet]:
    return [Sheet("Siteler", SITE_HEADER, [site_cells(row) for row in rows])]


def site_sheets(data: Mapping[str, Any]) -> list[Sheet]:
    summary = Sheet("Özet", ["Alan", "Değer"], period_rows(data))
    summary.rows += [
        list(pair) for pair in zip(SITE_HEADER, site_cells(data["totals"]), strict=True)
    ]
    return [
        Sheet(
            "Kullanıcılar",
            SITE_USER_HEADER,
            [user_cells(row, site=True) for row in data["users"]],
        ),
        summary,
        Sheet(
            "Ana bilgisayarlar",
            ["Ana bilgisayar", "Süre (sn)", "Bayt", "Oturum"],
            [[h["hostname"], h["seconds"], h["bytes"], h["sessions"]] for h in data["hostnames"]],
        ),
    ]


def categories_sheets(rows: Sequence[Mapping[str, Any]]) -> list[Sheet]:
    return [Sheet("Kategoriler", CATEGORY_HEADER, [category_cells(row) for row in rows])]


def category_sheets(data: Mapping[str, Any]) -> list[Sheet]:
    summary = Sheet("Özet", ["Alan", "Değer"], period_rows(data))
    summary.rows += [
        list(pair) for pair in zip(CATEGORY_HEADER, category_cells(data["totals"]), strict=True)
    ]
    return [
        Sheet("Siteler", SITE_HEADER, [site_cells(row) for row in data["sites"]]),
        summary,
        Sheet(
            "Kullanıcılar",
            CATEGORY_USER_HEADER,
            [user_cells(row, site=False) for row in data["users"]],
        ),
    ]


# -- files ----------------------------------------------------------------------------------------


def _csv_value(value: Any) -> Any:
    if isinstance(value, float):
        return f"{value:.4f}".replace(".", ",")
    return value


def csv_file(sheets: Sequence[Sheet], filename: str) -> Response:
    """The first sheet as CSV."""
    main = sheets[0]
    rows = ([_csv_value(cell) for cell in row] for row in main.rows)
    return csv_response(write_csv(main.header, rows), filename)


def _cell(sheet: Any, value: Any, *, bold: bool = False) -> Any:
    if isinstance(value, datetime):
        value = to_site(value).replace(tzinfo=None)
    elif isinstance(value, bool):
        value = "evet" if value else "hayır"
    elif isinstance(value, date):
        value = value.isoformat()
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub(lambda found: f"\\x{ord(found.group()):02x}", value)
    cell = WriteOnlyCell(sheet, value=value)
    if isinstance(value, str):
        cell.data_type = "s"  # never a formula: host names and URLs come from the network
    elif isinstance(value, datetime):
        cell.number_format = "dd.mm.yyyy hh:mm:ss"
    if bold:
        cell.font = Font(bold=True)
    return cell


def _write(sheet: Any, rows: Iterable[Sequence[Any]], *, bold: bool) -> None:
    for row in rows:
        sheet.append([_cell(sheet, value, bold=bold) for value in row])


def xlsx_bytes(sheets: Sequence[Sheet]) -> bytes:
    book = Workbook(write_only=True)
    for item in sheets:
        sheet = book.create_sheet(item.name[:SHEET_NAME_MAX])
        _write(sheet, [item.header], bold=True)
        _write(sheet, item.rows, bold=False)
    note = book.create_sheet("Yöntem")
    _write(note, [[METHOD_NOTE[0]], [METHOD_NOTE[1]]], bold=False)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def xlsx_file(sheets: Sequence[Sheet], filename: str) -> Response:
    return Response(
        content=xlsx_bytes(sheets),
        media_type=XLSX_TYPE,
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Cache-Control": "no-store",
        },
    )
