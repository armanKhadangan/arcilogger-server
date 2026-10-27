"""The rows of a search as a CSV or Excel file (contract 2.12, R17, R18).

* **Rows:** the search's own plan read by the window scan in keyset chunks (one read-only
  snapshot) until ``search.export_max_rows`` (100,000) or ``search.export_budget_seconds``
  (120 s); then the file stops and says so — in its name (``-kismi``), in the "Bilgi" sheet, in
  the header ``X-Arci-Export-Truncated`` (``rows`` or ``time``) and in the audit entry. Never
  silent, never a count first. The fill-in queries of a chunk (typed fields, ``messages`` text)
  get the time the budget has left; one that times out cuts the file there (``time``), so no
  row is written with empty columns.
* **CSV:** the rules of :mod:`arcilogger_api.csvfiles` (UTF-8 with BOM, ``;``, CRLF, every cell
  neutralised against formula injection), written row by row.
* **XLSX:** openpyxl write-only, every cell text (as ``usage_exports``: control characters an
  Excel file cannot hold written visibly as ``\\xNN``); the sheet "Sonuçlar" and the sheet
  "Bilgi" (query, range, user, time, rows, why it stopped, and that the file is not sealed).
* The file is built in a ``SpooledTemporaryFile`` (16 MB in memory, then on disk), hashed,
  audited (``search.export`` with its SHA-256 and size) and committed **before** it is sent,
  with ``Repr-Digest: sha-256=:…:`` (RFC 9530, checked by the browser) and the name
  ``arama-<from>-<to>[-kismi].<csv|xlsx>`` (site time).

Times are site time with microseconds; texts display-safe (``windows.display``), ``Metin`` the
whole ``messages.text`` (at most 2,000 characters). The file is not evidence: an evidence
package (Kanıt Paketi) carries the sealed raw lines.
"""

from __future__ import annotations

import base64
import csv
import hashlib
import io
import tempfile
from collections.abc import Iterator, Sequence
from dataclasses import dataclass
from datetime import datetime
from time import monotonic
from typing import IO, Any, Final, Literal

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from sqlalchemy.orm import Session

from arcilogger_api.csvfiles import BOM, SEPARATOR, neutralise
from arcilogger_api.search.cursor import Cursor
from arcilogger_api.search.fields import SiteNetwork
from arcilogger_api.search.planner import Plan
from arcilogger_api.search.windows import (
    Clock,
    SearchTimeoutError,
    WindowHook,
    WindowTimeoutError,
    page_rows,
    run_page,
)
from arcilogger_api.usage_exports import XLSX_TYPE
from arcilogger_core.timeutil import isoformat_site, to_site
from arcilogger_db.kinds import KINDS

Format = Literal["csv", "xlsx"]
Truncated = Literal["rows", "time"]
#: Rows read per page while exporting.
CHUNK: Final = 1000
#: ``messages.text`` is at most 2,000 characters (``parsers.base``).
TEXT_MAX: Final = 2000
#: Kept in memory up to this size, then in a temporary file.
SPOOL_BYTES: Final = 16 * 1024 * 1024
READ_BYTES: Final = 1024 * 1024
MEDIA_TYPES: Final[dict[str, str]] = {"csv": "text/csv; charset=utf-8", "xlsx": XLSX_TYPE}
HEADER: Final = (
    "Zaman",
    "Alınma zamanı",
    "Kaynak",
    "Gönderici",
    "Tür",
    "Alt tür",
    "Kaynak IP",
    "Kaynak port",
    "Hedef IP",
    "Hedef port",
    "NAT IP",
    "NAT port",
    "Eylem",
    "Alan adı",
    "URL",
    "Kullanıcı",
    "Uygulama",
    "Kategori",
    "MAC",
    "Metin",
    "Ham kayıt",
)
NOT_SEALED: Final = (
    "Bu dosya mühürlü değildir; bütünlük için Kanıt Paketi'ni kullanın.",
    "This file is not sealed; use an evidence package for integrity.",
)
TRUNCATION_TEXTS: Final = {
    None: ("Tamamı (kesilmedi)", "Complete (not cut)"),
    "rows": ("Satır sınırında kesildi", "Cut at the row limit"),
    "time": ("Süre sınırında kesildi", "Cut at the time limit"),
}
SUMMARY_COLUMNS: Final = (
    "srcip",
    "srcport",
    "dstip",
    "dstport",
    "natip",
    "natport",
    "action",
    "domain",
    "url",
    "user",
    "app",
    "category",
    "mac",
)


@dataclass(frozen=True)
class Info:
    """What the "Bilgi" sheet says about the export."""

    query: str
    start: datetime
    end: datetime
    username: str
    created_at: datetime


@dataclass
class ExportFile:
    file: IO[bytes]
    format: Format
    rows: int
    truncated: Truncated | None
    size: int
    sha256: str
    filename: str

    @property
    def media_type(self) -> str:
        return MEDIA_TYPES[self.format]

    @property
    def repr_digest(self) -> str:
        """``Repr-Digest`` (RFC 9530) of the file."""
        digest = base64.b64encode(bytes.fromhex(self.sha256)).decode("ascii")
        return f"sha-256=:{digest}:"

    def chunks(self) -> Iterator[bytes]:
        self.file.seek(0)
        while data := self.file.read(READ_BYTES):
            yield data

    def close(self) -> None:
        self.file.close()


def _time(value: datetime | None) -> str:
    return "" if value is None else isoformat_site(value)


def cells(row: dict[str, Any]) -> list[str]:
    """One row of the file (every value as text)."""
    summary = row["summary"]
    kind = KINDS.get(row["kind"])
    raw = row["raw"]
    values: list[Any] = [
        _time(row["time"]),
        _time(row["time_received"]),
        row["source"],
        row["sender"],
        kind.label_tr if kind else row["kind"],
        row["subtype"],
        *(summary.get(name) for name in SUMMARY_COLUMNS),
        row["text"],
        f"{raw['source']}/{raw['day'].isoformat()}/{raw['line']}/{raw['offset']}",
    ]
    return ["" if value is None else str(value) for value in values]


@dataclass
class Reading:
    """Rows read so far and why the reading stopped."""

    count: int = 0
    truncated: Truncated | None = None


def read_rows(
    session: Session,
    plan: Plan,
    start: datetime,
    end: datetime,
    site: SiteNetwork,
    reading: Reading,
    *,
    max_rows: int,
    budget_seconds: float,
    clock: Clock = monotonic,
    hook: WindowHook | None = None,
) -> Iterator[list[str]]:
    """The file's rows, page by page; ``reading`` records the count and the truncation."""
    deadline = clock() + budget_seconds

    def remaining() -> float:
        return deadline - clock()

    cursor: Cursor | None = None
    while True:
        left = remaining()
        if left <= 0:
            reading.truncated = "time"
            return
        want = min(CHUNK, max_rows + 1 - reading.count)
        try:
            page = run_page(session, plan, start, end, cursor, want, left, clock=clock, hook=hook)
        except SearchTimeoutError:
            reading.truncated = "time"
            return
        rows = page.rows
        more = reading.count + len(rows) > max_rows
        if more:
            rows = rows[: max_rows - reading.count]
        try:
            filled = page_rows(session, plan, rows, site, text_max=TEXT_MAX, seconds=remaining)
        except WindowTimeoutError:  # this chunk is left out whole, never half filled
            reading.truncated = "time"
            return
        for row in filled:
            reading.count += 1
            yield cells(row)
        if more:
            reading.truncated = "rows"
            return
        if page.next_cursor is None:
            return
        cursor = page.next_cursor


def _write_csv(target: IO[bytes], rows: Iterator[Sequence[str]]) -> None:
    text = io.TextIOWrapper(target, encoding="utf-8", newline="")
    try:
        text.write(BOM)
        writer = csv.writer(text, delimiter=SEPARATOR, lineterminator="\r\n")
        writer.writerow([neutralise(name) for name in HEADER])
        for row in rows:
            writer.writerow([neutralise(value) for value in row])
        text.flush()
    finally:
        text.detach()  # the caller keeps the file


def _cell(sheet: Any, value: str, *, bold: bool = False) -> Any:
    shown = ILLEGAL_CHARACTERS_RE.sub(lambda found: f"\\x{ord(found.group()):02x}", value)
    cell = WriteOnlyCell(sheet, value=shown)
    cell.data_type = "s"  # never a formula: host names and URLs come from the network
    if bold:
        cell.font = Font(bold=True)
    return cell


def info_rows(info: Info, reading: Reading) -> list[tuple[str, str]]:
    """The "Bilgi" sheet."""
    truncated_tr, truncated_en = TRUNCATION_TEXTS[reading.truncated]
    return [
        ("Sorgu / Query", info.query),
        ("Başlangıç / From", isoformat_site(info.start)),
        ("Bitiş / To", isoformat_site(info.end)),
        ("Kullanıcı / User", info.username),
        ("Oluşturulma / Created", isoformat_site(info.created_at)),
        ("Satır / Rows", str(reading.count)),
        ("Durum / Status", f"{truncated_tr} / {truncated_en}"),
        ("Not", NOT_SEALED[0]),
        ("Note", NOT_SEALED[1]),
    ]


def _write_xlsx(
    target: IO[bytes], rows: Iterator[Sequence[str]], info: Info, reading: Reading
) -> None:
    book = Workbook(write_only=True)
    results = book.create_sheet("Sonuçlar")
    results.append([_cell(results, name, bold=True) for name in HEADER])
    for row in rows:
        results.append([_cell(results, value) for value in row])
    sheet = book.create_sheet("Bilgi")
    for label, value in info_rows(info, reading):
        sheet.append([_cell(sheet, label, bold=True), _cell(sheet, value)])
    book.save(target)


def filename(start: datetime, end: datetime, extension: str, *, partial: bool) -> str:
    """``arama-<from>-<to>[-kismi].<ext>`` in site time."""
    first = to_site(start).strftime("%Y%m%dT%H%M")
    last = to_site(end).strftime("%Y%m%dT%H%M")
    return f"arama-{first}-{last}{'-kismi' if partial else ''}.{extension}"


def build(rows: Iterator[Sequence[str]], fmt: Format, info: Info, reading: Reading) -> ExportFile:
    """Write the file of ``rows`` (``reading`` is filled while they are read) and hash it."""
    target = tempfile.SpooledTemporaryFile(max_size=SPOOL_BYTES)  # noqa: SIM115 - returned
    try:
        if fmt == "csv":
            _write_csv(target, rows)
        else:
            _write_xlsx(target, rows, info, reading)
        target.seek(0)
        digest = hashlib.sha256()
        size = 0
        while data := target.read(READ_BYTES):
            digest.update(data)
            size += len(data)
    except BaseException:
        target.close()
        raise
    return ExportFile(
        file=target,
        format=fmt,
        rows=reading.count,
        truncated=reading.truncated,
        size=size,
        sha256=digest.hexdigest(),
        filename=filename(info.start, info.end, fmt, partial=reading.truncated is not None),
    )
